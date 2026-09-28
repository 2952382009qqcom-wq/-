"""明鉴 - 基于大模型的法律文书智能助手"""
import json
import io
import os
import re
import secrets
from functools import wraps
from flask import Flask, render_template, request, jsonify, send_file, Response, stream_with_context
from flask_cors import CORS
from flask_login import LoginManager, login_required, current_user
from fpdf import FPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from werkzeug.security import generate_password_hash
from sqlalchemy import inspect, text

from core.llm_client import call_llm_json, call_llm_stream_json, update_config, is_demo_mode, test_llm_connection
from core.prompts import (
    SYSTEM_PROMPT,
    ANALYZE_PROMPT,
    SEARCH_PROVISION_PROMPT,
    REVIEW_CONTRACT_PROMPT,
    GENERATE_DOCUMENT_PROMPT,
    STRATEGY_PROMPT,
    STUDENT_LEGAL_PROMPT,
)
from core.doc_parser import (
    DocumentParseError,
    ocr_available,
    parse_upload_detailed,
)
from core.privacy import (
    RedactionSession,
    public_redaction_summary,
    redact_nested_json,
)
from core.models import db, User, AnalysisRecord
from core.extensions import limiter, migrate, socketio
from core.community.routes import community_bp
from core.cases.routes import cases_bp
from core.chat.routes import chat_bp
from core.notifications.routes import notifications_bp
from core.memory.routes import memory_bp
from core.cases.seed import seed_reference_data
from core.recommendations.events import record_event
from core.taxonomy import infer_legal_domain
from core.legal_guidance import legal_anchors_for_question, requires_legal_basis
# Import model modules before db.create_all so every new table is registered.
from core.community import models as _community_models  # noqa: F401
from core.cases import models as _case_models  # noqa: F401
from core.chat import models as _chat_models  # noqa: F401
from core.recommendations import models as _recommendation_models  # noqa: F401
from core.notifications import models as _notification_models  # noqa: F401
from core.memory import models as _memory_models  # noqa: F401
from core.memory.services import build_memory_context
from core.conversation import (
    ConversationNotFound,
    add_message,
    conversation_context,
    delete_conversation,
    get_or_create_conversation,
    list_conversations,
    load_state,
    save_state,
    serialize_conversation,
)
from core.legal_agent import (
    INTENT_ANALYZE,
    INTENT_DRAFT,
    INTENT_EVIDENCE,
    INTENT_GENERAL,
    INTENT_REVIEW,
    INTENT_RISK,
    INTENT_SEARCH,
    INTENT_STRATEGY,
    build_general_prompt,
    compose_chat_answer,
    detect_intent,
    document_follow_up,
    extract_document_fields,
    infer_document_type,
    missing_document_fields,
    suggested_follow_ups,
)
from core.auth import auth_bp

app = Flask(__name__)
cors_origins = [
    origin.strip()
    for origin in os.getenv("CORS_ORIGINS", "http://localhost:5000,http://127.0.0.1:5000").split(",")
    if origin.strip()
]
CORS(app, supports_credentials=True, origins=cors_origins)

app.secret_key = os.getenv("SECRET_KEY", secrets.token_hex(32))
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "0").strip().lower() in {"1", "true", "yes"},
)
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///chatlaw.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
try:
    app.config["MAX_CONTENT_LENGTH"] = max(1, int(float(os.getenv("MAX_UPLOAD_MB", "15")) * 1024 * 1024))
except (TypeError, ValueError):
    app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024
db.init_app(app)
migrate.init_app(app, db)
app.config["RATELIMIT_STORAGE_URI"] = os.getenv("RATELIMIT_STORAGE_URI", "memory://")
limiter.init_app(app)
socketio.init_app(
    app,
    cors_allowed_origins=cors_origins,
    message_queue=os.getenv("SOCKETIO_REDIS_URL", "").strip() or None,
    logger=False,
    engineio_logger=False,
)

login_manager = LoginManager()
login_manager.init_app(app)

@login_manager.unauthorized_handler
def unauthorized():
    return jsonify({"error": "请先登录"}), 401


@app.before_request
def reject_cross_site_mutations():
    """CSRF boundary for session-authenticated JSON/upload APIs and Socket handshakes."""
    if request.method in {"GET", "HEAD", "OPTIONS"}:
        return None
    if request.headers.get("Sec-Fetch-Site", "").lower() == "cross-site":
        return jsonify({"error": "拒绝跨站请求", "code": "csrf_rejected"}), 403
    origin = request.headers.get("Origin", "").rstrip("/")
    allowed = {value.rstrip("/") for value in cors_origins}
    if origin and origin not in allowed:
        return jsonify({"error": "请求来源不受信任", "code": "csrf_rejected"}), 403
    return None

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

app.register_blueprint(auth_bp)
app.register_blueprint(community_bp)
app.register_blueprint(cases_bp)
app.register_blueprint(chat_bp)
app.register_blueprint(notifications_bp)
app.register_blueprint(memory_bp)
# Event decorators are registered as a side effect of this import.
from core.chat import socket_events as _chat_socket_events  # noqa: E402,F401

APP_VERSION = os.getenv("APP_VERSION", "2026.09-legal-agent")


def _text_field(data, key, default=""):
    """Return a stripped JSON text field, or None for an invalid type."""
    value = data.get(key, default) if isinstance(data, dict) else default
    return value.strip() if isinstance(value, str) else None


@app.errorhandler(DocumentParseError)
def handle_document_parse_error(error):
    return jsonify({"error": error.message, "code": error.code}), error.status_code


@app.errorhandler(413)
def handle_request_too_large(_error):
    return jsonify({
        "error": "上传文件超过服务器允许的大小。",
        "code": "file_too_large",
    }), 413


def ensure_user_schema():
    inspector = inspect(db.engine)
    columns = {column["name"] for column in inspector.get_columns("users")}
    if "approval_requested_at" not in columns:
        db.session.execute(text("ALTER TABLE users ADD COLUMN approval_requested_at DATETIME"))
    if "llm_api_key" not in columns:
        db.session.execute(text("ALTER TABLE users ADD COLUMN llm_api_key VARCHAR(512)"))
    if "llm_base_url" not in columns:
        db.session.execute(text("ALTER TABLE users ADD COLUMN llm_base_url VARCHAR(255)"))
    if "free_api_used" not in columns:
        db.session.execute(text("ALTER TABLE users ADD COLUMN free_api_used INTEGER NOT NULL DEFAULT 0"))
    db.session.commit()


def ensure_application_schema():
    """Create additive tables for both direct and WSGI-based production starts."""
    with app.app_context():
        db.create_all()
        ensure_user_schema()
        if os.getenv("AUTO_SEED_REFERENCE_DATA", "1").strip().lower() not in {"0", "false", "no"}:
            seed_reference_data()


if os.getenv("AUTO_INIT_DB", "1").strip().lower() not in {"0", "false", "no"}:
    ensure_application_schema()


def _prepare_document_for_model(text, max_chars, parsed=None, filename=""):
    """Redact PII before any external model call and expose safe pipeline metadata."""
    session = RedactionSession()
    safe_text = session.redact(text or "")
    privacy_meta = public_redaction_summary(session)
    privacy_meta["applied_before_model"] = True

    if parsed is not None:
        document_meta = dict(parsed.metadata)
    else:
        document_meta = {
            "method": "text_input",
            "page_count": 1,
            "char_count": len(text or ""),
            "warnings": [],
            "ocr_confidence": None,
        }
    if filename:
        document_meta["filename"] = os.path.basename(filename)

    document_meta["truncated"] = len(safe_text) > max_chars
    if document_meta["truncated"]:
        document_meta.setdefault("warnings", []).append(
            f"文本共 {len(safe_text)} 字，本次模型分析使用前 {max_chars} 字。"
        )
        safe_text = safe_text[:max_chars]
    document_meta["analyzed_char_count"] = len(safe_text)
    return safe_text, privacy_meta, document_meta


def _attach_processing_meta(result, privacy_meta, document_meta, final_stage="analyze"):
    result["privacy_meta"] = privacy_meta
    result["document_meta"] = document_meta
    result["pipeline_trace"] = [
        {
            "stage": "extract",
            "status": "ok",
            "method": document_meta.get("method", "text_input"),
            "page_count": document_meta.get("page_count", 1),
        },
        {
            "stage": "privacy",
            "status": "ok",
            "masked_count": privacy_meta.get("masked_count", 0),
        },
        {"stage": final_stage, "status": "ok"},
    ]
    return result


def _free_api_limit():
    try:
        return max(0, int(os.getenv("FREE_API_LIMIT", "10")))
    except ValueError:
        return 10


def _free_api_remaining():
    used = max(0, int(getattr(current_user, "free_api_used", 0) or 0))
    return max(0, _free_api_limit() - used)


def _refresh_current_user():
    db.session.expire(current_user._get_current_object(), ["free_api_used"])


def _refund_free_api_call(user_id=None):
    target_user_id = user_id
    if target_user_id is None:
        target_user_id = current_user.id
    db.session.execute(
        text(
            "UPDATE users SET free_api_used = free_api_used - 1 "
            "WHERE id = :user_id AND free_api_used > 0"
        ),
        {"user_id": target_user_id},
    )
    db.session.commit()
    if user_id is None:
        _refresh_current_user()


def require_approved(f):
    @wraps(f)
    def decorated(*a, **kw):
        if not current_user.is_authenticated:
            return jsonify({"error": "请先登录"}), 401
        has_personal_api = bool(getattr(current_user, "llm_api_key", None))
        has_unlimited_access = current_user.is_approved or current_user.is_admin or has_personal_api
        if has_unlimited_access:
            return f(*a, **kw)

        limit = _free_api_limit()
        reservation = db.session.execute(
            text(
                "UPDATE users SET free_api_used = COALESCE(free_api_used, 0) + 1 "
                "WHERE id = :user_id AND COALESCE(free_api_used, 0) < :limit"
            ),
            {"user_id": current_user.id, "limit": limit},
        )
        db.session.commit()
        _refresh_current_user()

        if reservation.rowcount != 1:
            return jsonify({
                "error": "10 次免费 AI 调用额度已用完，请在 API 设置中配置自己的 API Key。",
                "code": "FREE_API_QUOTA_EXHAUSTED",
                "free_api_limit": limit,
                "free_api_remaining": 0,
            }), 403

        try:
            response = app.make_response(f(*a, **kw))
        except Exception:
            _refund_free_api_call()
            raise

        if response.status_code >= 400:
            _refund_free_api_call()
        elif response.mimetype == "text/event-stream":
            original_iterable = response.response
            reserved_user_id = current_user.id

            def guarded_stream():
                try:
                    yield from original_iterable
                except Exception:
                    app.logger.exception("Streaming AI response failed")
                    try:
                        with app.app_context():
                            _refund_free_api_call(reserved_user_id)
                    except Exception:
                        app.logger.exception("Failed to refund interrupted free API call")
                    payload = {
                        "done": True,
                        "result": {"error": "流式响应异常中断，本次免费额度已退回"},
                    }
                    yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

            response.response = guarded_stream()
        response.headers["X-Free-Api-Limit"] = str(limit)
        response.headers["X-Free-Api-Remaining"] = str(_free_api_remaining())
        return response
    return decorated


@app.after_request
def configure_streaming_headers(response):
    if response.mimetype == "text/event-stream":
        response.headers.setdefault("Cache-Control", "no-cache, no-transform")
        response.headers.setdefault("X-Accel-Buffering", "no")
    return response


def save_record(module_type, input_text, result):
    if current_user.is_authenticated:
        try:
            redaction_session = RedactionSession()
            safe_input = redaction_session.redact(input_text or "")
            safe_result = redact_nested_json(result, redaction_session)
            record = AnalysisRecord(
                user_id=current_user.id,
                module_type=module_type,
                input_text=safe_input[:5000],
                result_json=json.dumps(safe_result, ensure_ascii=False)[:10000],
            )
            db.session.add(record)
            db.session.commit()
            extra_records = (
                AnalysisRecord.query
                .filter_by(user_id=current_user.id)
                .order_by(AnalysisRecord.created_at.desc(), AnalysisRecord.id.desc())
                .offset(20)
                .all()
            )
            for old_record in extra_records:
                db.session.delete(old_record)
            if extra_records:
                db.session.commit()
        except Exception:
            db.session.rollback()


# ===== 首页 =====
@app.route("/")
def index():
    """渲染主页面"""
    return render_template("index.html")


# ===== 演示数据生成（无 API Key 时自动启用） =====

def _demo_analyze(text):
    return {
        "demo_mode": True,
        "document_type": "买卖合同",
        "parties": ["甲方（买受人）", "乙方（出卖人）"],
        "key_clauses": [
            {"clause": "交货条款", "summary": "约定交货时间和地点，但未明确验收标准", "risk_level": "中", "note": "建议补充验收标准"},
            {"clause": "付款条款", "summary": "付款节点与交货进度挂钩", "risk_level": "低", "note": ""},
            {"clause": "违约责任条款", "summary": "仅约定了乙方违约责任，甲方违约责任缺失", "risk_level": "高", "note": "违约责任不对等"},
        ],
        "legal_basis": [],
        "risk_points": [
            {"point": "违约责任不对等", "level": "高", "suggestion": "应约定双方对等的违约责任条款"},
            {"point": "验收标准不明确", "level": "中", "suggestion": "明确验收时间、标准和异议提出方式"},
            {"point": "知识产权归属未约定", "level": "中", "suggestion": "如涉及技术成果，应明确知识产权归属"},
        ],
        "overall_assessment": "该合同整体框架完整，但存在违约责任不对等、关键条款缺失等问题，建议在签署前进行修订完善。",
        "revision_suggestions": [
            "增加双方对等的违约责任条款",
            "明确交货验收标准和验收期限",
            "补充知识产权归属及保密条款",
            "明确争议解决方式和管辖法院",
        ],
    }


def _demo_search(question):
    return {
        "demo_mode": True,
        "question": question,
        "provisions": [],
        "answer": "当前未连接可用的外部模型，无法可靠回答这个法律问题。请稍后重试或先配置模型 API。",
        "legal_analysis": "当前未连接可用的外部模型，无法可靠生成法律分析。",
        "practical_advice": "请补充事情经过、发生时间、所在地区、对方身份和已有证据后重试。",
        "related_cases": [],
        "evidence_checklist": [],
        "risk_points": [],
        "next_steps": [],
        "questions_to_clarify": ["事情发生在什么时间和地区？", "目前有哪些合同、转账和聊天记录？"],
    }


def _demo_review(contract):
    return {
        "demo_mode": True,
        "contract_type": "货物买卖合同",
        "overall_risk_score": 72,
        "overall_risk_level": "中",
        "risk_items": [
            {"clause_text": "如乙方逾期交货，每逾期一日按合同金额0.1%支付违约金...", "risk_type": "违约责任不对等",
             "risk_level": "高", "risk_score": 85, "explanation": "仅约定了乙方违约责任，未约定甲方逾期付款的违约责任，违反公平原则。",
             "revised_text": "双方应对等约定：甲方逾期付款的，每逾期一日按未付金额0.1%支付违约金；乙方逾期交货的，每逾期一日按合同金额0.1%支付违约金。",
             "legal_basis": ""},
            {"clause_text": "因履行本合同发生争议，由乙方所在地人民法院管辖。", "risk_type": "争议解决条款不利",
             "risk_level": "中", "risk_score": 70, "explanation": "约定对方所在地法院管辖，增加己方维权成本。",
             "revised_text": "因履行本合同发生争议，由合同签订地（甲方所在地）人民法院管辖。", "legal_basis": ""},
            {"clause_text": "乙方有权在提前通知甲方后单方解除本合同...", "risk_type": "单方解除权",
             "risk_level": "高", "risk_score": 80, "explanation": "仅约定一方解除权，未约定解除后的清算事宜。",
             "revised_text": "双方可在符合法律规定或约定条件时解除合同，并约定解除后的清算方式。",
             "legal_basis": ""},
        ],
        "missing_clauses": ["知识产权归属条款", "保密条款", "不可抗力条款", "送达条款"],
        "summary": "该合同整体风险为中等偏高，主要存在违约责任不对等、争议解决管辖不利等问题。建议在签署前对高风险条款进行修订，并补充缺失的必要条款。",
    }


COMPLAINT_FIELD_LABELS = {
    "plaintiff": "原告",
    "plaintiff_address": "住所",
    "legal_representative": "法定代表人/主要负责人",
    "representative_duty": "职务",
    "representative_contact": "联系方式",
    "agent": "委托诉讼代理人",
    "defendant": "被告",
    "defendant_info": "被告基本信息",
    "claims": "诉讼请求",
    "facts_and_reasons": "事实和理由",
    "evidence": "证据和证据来源，证人姓名和住所",
    "court": "受诉人民法院",
    "copy_count": "副本份数",
    "suer": "起诉人",
    "date": "日期",
}


APPEAL_FIELD_LABELS = {
    "appellant": "上诉人(原审诉讼地位)",
    "appellant_info": "上诉人基本信息",
    "appellant_contact": "联系方式",
    "legal_agent": "法定代理人/指定代理人",
    "agent": "委托诉讼代理人",
    "appellee": "被上诉人(原审诉讼地位)",
    "appellee_info": "被上诉人基本信息",
    "case_parties": "案由当事人",
    "cause": "案由",
    "original_court": "原审人民法院",
    "judgment_date": "原审裁判日期",
    "case_number": "原审裁判文号",
    "judgment_type": "裁判类型",
    "appeal_requests": "上诉请求",
    "appeal_reasons": "上诉理由",
    "court": "受诉人民法院",
    "copy_count": "副本份数",
    "date": "日期",
}


DEFENSE_FIELD_LABELS = {
    "respondent": "答辩人",
    "respondent_info": "答辩人基本信息",
    "respondent_contact": "联系方式",
    "legal_agent": "法定代理人/指定代理人",
    "agent": "委托诉讼代理人",
    "original_court": "受理人民法院",
    "case_number": "案号",
    "case_summary": "当事人和案由",
    "defense_opinion": "答辩意见",
    "evidence": "证据和证据来源，证人姓名和住所",
    "court": "受诉人民法院",
    "copy_count": "副本份数",
    "date": "日期",
}


EXECUTION_FIELD_LABELS = {
    "applicant": "申请执行人",
    "applicant_info": "申请执行人基本信息",
    "applicant_contact": "联系方式",
    "legal_agent": "法定代理人/指定代理人",
    "agent": "委托诉讼代理人",
    "respondent": "被执行人",
    "respondent_info": "被执行人基本信息",
    "case_parties": "执行案件当事人",
    "cause": "案由",
    "instrument_maker": "生效法律文书作出机关",
    "instrument_number": "生效法律文书文号",
    "instrument_type": "生效法律文书类型",
    "obligor": "未履行义务人",
    "execution_requests": "请求事项",
    "court": "受诉人民法院",
    "attachment_count": "生效法律文书份数",
    "date": "日期",
}


COUNTERCLAIM_FIELD_LABELS = {
    "counter_plaintiff": "反诉原告(本诉被告)",
    "counter_plaintiff_info": "反诉原告基本信息",
    "counter_plaintiff_contact": "联系方式",
    "legal_agent": "法定代理人/指定代理人",
    "agent": "委托诉讼代理人",
    "counter_defendant": "反诉被告(本诉原告)",
    "counter_defendant_info": "反诉被告基本信息",
    "counterclaim_requests": "反诉请求",
    "facts_and_reasons": "事实和理由",
    "evidence": "证据和证据来源，证人姓名和住所",
    "court": "受诉人民法院",
    "copy_count": "副本份数",
    "date": "日期",
}


JURISDICTION_OBJECTION_FIELD_LABELS = {
    "objector": "异议人(被告)",
    "objector_info": "异议人基本信息",
    "objector_contact": "联系方式",
    "legal_agent": "法定代理人/指定代理人",
    "agent": "委托诉讼代理人",
    "original_court": "原受理人民法院",
    "case_number": "案号",
    "case_summary": "案件当事人和案由",
    "transfer_court": "移送人民法院",
    "facts_and_reasons": "事实和理由",
    "court": "受诉人民法院",
    "date": "日期",
}


JUDICIAL_CONFIRMATION_FIELD_LABELS = {
    "applicant_person": "自然人申请人",
    "applicant_person_info": "自然人申请人基本信息",
    "person_agent_1": "自然人代理人一",
    "person_agent_2": "自然人代理人二",
    "applicant_org": "单位申请人",
    "applicant_org_info": "单位申请人基本信息",
    "legal_representative": "法定代表人/主要负责人",
    "org_agent_1": "单位代理人一",
    "org_agent_2": "单位代理人二",
    "claims": "诉讼请求",
    "facts_and_reasons": "事实和理由",
    "court": "受诉人民法院",
    "date": "日期",
}


CITIZEN_AUTHORIZATION_FIELD_LABELS = {
    "principal": "委托人",
    "principal_info": "委托人基本信息",
    "lawyer_agent": "律师受委托人",
    "citizen_agent": "公民受委托人",
    "case_summary": "当事人和案由",
    "agent_names": "委托代理人姓名",
    "agent_1_name": "代理人一姓名",
    "agent_1_authority": "代理人一事项和权限",
    "agent_2_name": "代理人二姓名",
    "agent_2_authority": "代理人二事项和权限",
    "date": "日期",
}


def _clean_complaint_field(value):
    return str(value or "").strip()


def _extract_fields_from_description(description, labels):
    fields = {key: "" for key in labels}
    for raw_line in str(description or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        for key, label in labels.items():
            if line.startswith(f"{label}：") or line.startswith(f"{label}:"):
                fields[key] = line.split("：", 1)[-1].strip() if "：" in line else line.split(":", 1)[-1].strip()
                break
    return fields


def _extract_complaint_fields_from_description(description):
    return _extract_fields_from_description(description, COMPLAINT_FIELD_LABELS)


def _extract_appeal_fields_from_description(description):
    return _extract_fields_from_description(description, APPEAL_FIELD_LABELS)


def _extract_defense_fields_from_description(description):
    return _extract_fields_from_description(description, DEFENSE_FIELD_LABELS)


def _extract_execution_fields_from_description(description):
    return _extract_fields_from_description(description, EXECUTION_FIELD_LABELS)


def _extract_counterclaim_fields_from_description(description):
    return _extract_fields_from_description(description, COUNTERCLAIM_FIELD_LABELS)


def _extract_jurisdiction_objection_fields_from_description(description):
    return _extract_fields_from_description(description, JURISDICTION_OBJECTION_FIELD_LABELS)


def _extract_judicial_confirmation_fields_from_description(description):
    return _extract_fields_from_description(description, JUDICIAL_CONFIRMATION_FIELD_LABELS)


def _extract_citizen_authorization_fields_from_description(description):
    return _extract_fields_from_description(description, CITIZEN_AUTHORIZATION_FIELD_LABELS)


def _normalize_multiline_items(value):
    text = _clean_complaint_field(value)
    if not text:
        return "……"
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) > 1:
        return "\n".join(lines)
    parts = [part.strip() for part in re.split(r"[；;]", text) if part.strip()]
    return "\n".join(parts) if len(parts) > 1 else text


def _build_civil_complaint_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    plaintiff = get("plaintiff")
    plaintiff_address = get("plaintiff_address")
    legal_representative = get("legal_representative")
    representative_duty = get("representative_duty")
    representative_contact = get("representative_contact")
    agent = get("agent")
    defendant = get("defendant")
    defendant_info = _clean_complaint_field(fields.get("defendant_info"))
    claims = _normalize_multiline_items(fields.get("claims"))
    facts_and_reasons = _normalize_multiline_items(fields.get("facts_and_reasons"))
    evidence = _normalize_multiline_items(fields.get("evidence"))
    court = get("court", "××××人民法院")
    copy_count = get("copy_count", "×")
    suer = get("suer", "起诉人")
    date = get("date", "××××年××月××日")

    representative_parts = []
    if legal_representative != "……":
        representative_parts.append(legal_representative)
    if representative_duty != "……":
        representative_parts.append(f"{representative_duty}(写明职务)")
    if representative_contact != "……":
        representative_parts.append(f"联系方式：{representative_contact}")
    representative_line = "，".join(representative_parts) if representative_parts else "×××，……(写明职务)，联系方式：……"
    defendant_line = defendant if not defendant_info else f"{defendant}，{defendant_info}"

    return "\n".join([
        "民事起诉状",
        "",
        f"原告：{plaintiff}，住所：{plaintiff_address}。",
        f"法定代表人/主要负责人：{representative_line}。",
        f"委托诉讼代理人：{agent}。",
        f"被告：{defendant_line}。",
        "……",
        "(以上写明当事人和其他诉讼参加人的姓名或者名称等基本信息)",
        "",
        "诉讼请求：",
        claims,
        "",
        "事实和理由：",
        facts_and_reasons,
        "",
        "证据和证据来源，证人姓名和住所：",
        evidence,
        "",
        "此致",
        court,
        "",
        f"附：本起诉状副本{copy_count}份",
        "",
        "起诉人(公章和签名)",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第一百二十条第一款、第一百二十一条制定，供法人或者其他组织提起民事诉讼用。",
        "2．起诉应当向人民法院递交起诉状，并按照被告人数提出副本。",
        "3．起诉时已经委托诉讼代理人的，应当写明委托诉讼代理人基本信息。",
        "4．被告是自然人的，应当写明姓名、性别、工作单位、住所等信息；被告是法人或者其他组织的，应当写明名称、住所等信息。",
        "5．原告在起诉状中直接列写第三人的，视为其申请人民法院追加该第三人参加诉讼。是否通知第三人参加诉讼，由人民法院审查决定。",
        "6．起诉状应当加盖单位印章，并由法定代表人或者主要负责人签名。",
    ])


def _build_civil_appeal_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    appellant = get("appellant", "×××")
    appellant_info = get("appellant_info", "男/女，××××年××月××日出生，×族，……(写明工作单位和职务或者职业)，住……")
    appellant_contact = get("appellant_contact")
    legal_agent = get("legal_agent")
    agent = get("agent")
    appellee = get("appellee", "×××")
    appellee_info = get("appellee_info")
    case_parties = get("case_parties", "×××因与×××")
    cause = get("cause", "……(写明案由)")
    original_court = get("original_court", "××××人民法院")
    judgment_date = get("judgment_date", "××××年××月××日")
    case_number = get("case_number", "(××××)……号")
    judgment_type = get("judgment_type", "民事判决/裁定")
    appeal_requests = _normalize_multiline_items(fields.get("appeal_requests"))
    appeal_reasons = _normalize_multiline_items(fields.get("appeal_reasons"))
    court = get("court", "××××人民法院")
    copy_count = get("copy_count", "×")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "民事上诉状",
        "",
        f"上诉人(原审诉讼地位)：{appellant}，{appellant_info}。联系方式：{appellant_contact}。",
        f"法定代理人/指定代理人：{legal_agent}。",
        f"委托诉讼代理人：{agent}。",
        f"被上诉人(原审诉讼地位){appellee}，{appellee_info}。",
        "……",
        "(以上写明当事人和其他诉讼参加人的姓名或者名称等基本信息)",
        f"{case_parties}{cause}一案，不服{original_court}{judgment_date}作出的{case_number}{judgment_type}，现提起上诉。",
        "",
        "上诉请求：",
        appeal_requests,
        "",
        "上诉理由：",
        appeal_reasons,
        "",
        "此致",
        court,
        "",
        f"附：本上诉状副本{copy_count}份",
        "",
        "上诉人(签名或者盖章)",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第一百六十四条、第一百六十五条、第一百六十六条、第二百六十九条制定，供不服第一审人民法院民事判决或者裁定的当事人，向上一级人民法院提起上诉用。",
        "2．当事人是法人或者其他组织的，写明名称住所。另起一行写明法定代表人、主要负责人及其姓名、职务、联系方式。",
        "3．当事人不服地方人民法院第一审判决的，有权在判决书送达之日起十五日内向上一级人民法院提起上诉。当事人不服地方人民法院第一审裁定的，有权在裁定书送达之日起十日内向上一级人民法院提起上诉。在中华人民共和国领域内没有住所的当事人，不服第一审人民法院判决、裁定的，有权在判决书、裁定书送达之日起三十日内提起上诉。",
        "4．上诉状的内容，应当包括当事人的姓名，法人的名称及其法定代表人的姓名或者其他组织的名称及其主要负责人的姓名；原审人民法院名称、案件的编号和案由；上诉的请求和理由。",
        "5．上诉状应当通过原审人民法院提出，并按照对方当事人或者代表人的人数提出副本。",
        "6．有新证据的，应当在上诉理由之后写明证据和证据来源，证人姓名和住所。",
    ])


def _build_civil_defense_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    respondent = get("respondent", "×××")
    respondent_info = get("respondent_info", "男/女，××××年××月××日生，×族，……(写明工作单位和职务或职业)，住……")
    respondent_contact = get("respondent_contact")
    legal_agent = get("legal_agent")
    agent = get("agent")
    original_court = get("original_court", "××××人民法院")
    case_number = get("case_number", "(××××)……民初……号")
    case_summary = get("case_summary", "……(写明当事人和案由)")
    defense_opinion = _normalize_multiline_items(fields.get("defense_opinion"))
    evidence = _normalize_multiline_items(fields.get("evidence"))
    court = get("court", "××××人民法院")
    copy_count = get("copy_count", "×")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "民事答辩状",
        "",
        f"答辩人：{respondent}，{respondent_info}。联系方式：{respondent_contact}。",
        f"法定代理人/指定代理人：{legal_agent}。",
        f"委托诉讼代理人：{agent}。",
        "(以上写明答辩人和其他诉讼参加人的姓名或者名称等基本信息)",
        f"对{original_court}{case_number}{case_summary}一案的起诉，答辩如下：",
        defense_opinion,
        "证据和证据来源，证人姓名和住所：",
        evidence,
        "此致",
        court,
        "",
        f"附：本答辩状副本{copy_count}份",
        "",
        "答辩人(签名)",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第一百二十五条制定，供公民对民事起诉提出答辩用。",
        "2．被告应当在收到起诉状副本之日起十五日内提出答辩状。被告在中华人民共和国领域内没有住所的，应当在收到起诉状副本后三十日内提出答辩状。被告申请延期答辩的，是否准许，由人民法院决定。",
        "3．答辩状应当记明被告的姓名、性别、出生日期、民族、工作单位、职业、住所、联系方式。",
        "4．答辩时已经委托诉讼代理人的，应当写明委托诉讼代理人基本信息。",
        "5．答辩状应当由本人签名。",
    ])


def _build_execution_application_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    applicant = get("applicant", "×××")
    applicant_info = get("applicant_info", "男/女，××××年××月××日出生，×族，……(写明工作单位和职务或者职业)，住……")
    applicant_contact = get("applicant_contact")
    legal_agent = get("legal_agent")
    agent = get("agent")
    respondent = get("respondent", "×××")
    respondent_info = get("respondent_info")
    case_parties = get("case_parties", f"申请执行人{applicant}与被执行人{respondent}")
    cause = get("cause", "……(写明案由)")
    instrument_maker = get("instrument_maker", "××××人民法院(或其他生效法律文书的作出机关)")
    instrument_number = get("instrument_number", "(××××)……号")
    instrument_type = get("instrument_type", "民事判决(或其他生效法律文书)")
    obligor = get("obligor", respondent)
    execution_requests = _normalize_multiline_items(fields.get("execution_requests"))
    court = get("court", "××××人民法院")
    attachment_count = get("attachment_count", "×")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "申请执行书",
        "",
        f"申请执行人：{applicant}，{applicant_info}。联系方式：{applicant_contact}。",
        f"法定代理人/指定代理人：{legal_agent}。",
        f"委托诉讼代理人：{agent}。",
        f"被执行人：{respondent}，{respondent_info}。",
        "……",
        "(以上写明申请执行人、被执行人和其他诉讼参加人的姓名或者名称等基本信息)",
        f"{case_parties}{cause}一案，{instrument_maker}{instrument_number}{instrument_type}已发生法律效力。被执行人{obligor}未履行/未全部履行生效法律文书确定的给付义务，特向你院申请强制执行。",
        "请求事项",
        execution_requests,
        "此致",
        court,
        f"附：生效法律文书{attachment_count}份",
        "申请执行人(签名或盖章)",
        date,
        "",
        "【说明】",
        "1．本文书样式根据《中华人民共和国民事诉讼法》第二百三十六条、第二百三十七条第一款、第二百三十八条第一款，《最高人民法院关于人民法院执行工作若干问题的规定(试行)》第18条、19条、20条、21条、22条、23条规定制定，供申请执行人向人民法院申请执行时用。",
        "2．当事人是法人或者其他组织的，写明名称住所。另起一行写明法定代表人、主要负责人及其姓名、职务、联系方式。",
        "3．申请执行人向人民法院申请强制执行的内容，必须为生效法律文书确定的给付义务。",
    ])


def _build_civil_counterclaim_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    counter_plaintiff = get("counter_plaintiff", "×××")
    counter_plaintiff_info = get("counter_plaintiff_info", "男/女，××××年××月××日生，×族，……(写明工作单位和职务或职业)，住……")
    counter_plaintiff_contact = get("counter_plaintiff_contact")
    legal_agent = get("legal_agent")
    agent = get("agent")
    counter_defendant = get("counter_defendant", "×××")
    counter_defendant_info = get("counter_defendant_info")
    counterclaim_requests = _normalize_multiline_items(fields.get("counterclaim_requests"))
    facts_and_reasons = _normalize_multiline_items(fields.get("facts_and_reasons"))
    evidence = _normalize_multiline_items(fields.get("evidence"))
    court = get("court", "××××人民法院")
    copy_count = get("copy_count", "×")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "民事反诉状",
        "",
        f"反诉原告(本诉被告)：{counter_plaintiff}，{counter_plaintiff_info}。联系方式：{counter_plaintiff_contact}。",
        f"法定代理人/指定代理人：{legal_agent}。",
        f"委托诉讼代理人：{agent}。",
        f"反诉被告(本诉原告)：{counter_defendant}，{counter_defendant_info}。",
        "……",
        "(以上写明当事人和其他诉讼参加人的姓名或者名称等基本信息)",
        "反诉请求：",
        counterclaim_requests,
        "事实和理由：",
        facts_and_reasons,
        "证据和证据来源，证人姓名和住所：",
        evidence,
        "此致",
        court,
        "",
        f"附：本反诉状副本{copy_count}份",
        "",
        "反诉人(签名)",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第五十一条、第一百二十条第一款、第一百二十一条制定，供公民提起民事反诉用。",
        "2．反诉应当向人民法院递交反诉状，并按照被反诉人数提出副本。",
        "3．反诉原告应当写明姓名、性别、出生日期、民族、职业、工作单位、住所、联系方式。反诉原告是无民事行为能力或者限制民事行为能力人的，应当写明法定代理人姓名、性别、出生日期、民族、职业、工作单位、住所、联系方式，在诉讼地位后括注与原告的关系。",
        "4．反诉时已经委托诉讼代理人的，应当写明委托诉讼代理人基本信息。",
        "5．被反诉被告是自然人的，应当写明姓名、性别、工作单位、住所等信息；反诉被告是法人或者其他组织的，应当写明名称、住所等信息。",
        "6．反诉状应当由本人签名。",
    ])


def _build_jurisdiction_objection_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    objector = get("objector", "×××")
    objector_info = get("objector_info", "男/女，××××年××月××日出生，×族，……(写明工作单位和职务或者职业)，住……")
    objector_contact = get("objector_contact")
    legal_agent = get("legal_agent")
    agent = get("agent")
    original_court = get("original_court", "××××人民法院")
    case_number = get("case_number", "(××××)……号")
    case_summary = get("case_summary", "……(写明案件当事人和案由)")
    transfer_court = get("transfer_court", "××××人民法院")
    facts_and_reasons = _normalize_multiline_items(fields.get("facts_and_reasons"))
    court = get("court", "××××人民法院")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "异议书",
        "",
        f"异议人(被告)：{objector}，{objector_info}。联系方式：{objector_contact}。",
        f"法定代理人/指定代理人：{legal_agent}。",
        f"委托诉讼代理人：{agent}。",
        "(以上写明异议人和其他诉讼参加人的姓名或者名称等基本信息)",
        "请求事项：",
        f"将{original_court}{case_number}{case_summary}一案移送{transfer_court}管辖。",
        "事实和理由：",
        facts_and_reasons,
        "此致",
        court,
        "",
        "异议人(签名或者盖章)",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第一百二十七条第一款制定，供当事人向第一审人民法院提出管辖权异议用。",
        "2．当事人是法人或者其他组织的，写明名称住所。另起一行写明法定代表人、主要负责人及其姓名、职务、联系方式。",
        "3．人民法院受理案件后，当事人对管辖权有异议的，应当在提交答辩状期间提出。",
    ])


def _build_judicial_confirmation_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    applicant_person = get("applicant_person", "{姓名}")
    applicant_person_info = get("applicant_person_info", "{性别}，{出生年/月/日}生，{民族}，{证件类型}：{证件号码}，工作单位：{工作单位}，职业：{专业技术人员}，住{住所地}，联系方式：{手机号码}")
    person_agent_1 = get("person_agent_1", "{姓名}，律所：{律所名称}。联系方式：{手机号码}")
    person_agent_2 = get("person_agent_2", "{姓名}，律所：{律所名称}。联系方式：{手机号码}")
    applicant_org = get("applicant_org", "{单位名称}")
    applicant_org_info = get("applicant_org_info", "住所：{单位住所地}。{证照类型}：{证照号码}")
    legal_representative = get("legal_representative", "{姓名}，{证件类型}：{证件号码}，职务：{职务}，联系方式：{手机号码}")
    org_agent_1 = get("org_agent_1", "{姓名}，律所：{律所名称}。联系方式：{手机号码}")
    org_agent_2 = get("org_agent_2", "{姓名}，律所：{律所名称}。联系方式：{手机号码}")
    claims = _normalize_multiline_items(fields.get("claims"))
    facts_and_reasons = _normalize_multiline_items(fields.get("facts_and_reasons"))
    court = get("court", "{法院名称}")
    date = get("date", "××××年××月××日")

    return "\n".join([
        "申请书",
        "",
        f"申请人：{applicant_person}，{applicant_person_info}。",
        f"代理人：{person_agent_1}。",
        f"代理人：{person_agent_2}。",
        f"申请人：{applicant_org}，{applicant_org_info}，",
        f"法定代表人/主要负责人：{legal_representative}。",
        f"代理人：{org_agent_1}。",
        f"代理人：{org_agent_2}。",
        "诉讼请求：",
        claims,
        "事实和理由：",
        facts_and_reasons,
        "申请人出于解决纠纷的目的自愿达成协议，没有恶意串通、规避法律的行为；如果因为该协议内容而给国家、集体或他人造成损害的，愿意承担相应的民事责任和其他法律责任。",
        "此致",
        "",
        court,
        "",
        "申请人(签名或者盖章)　　",
        date,
    ])


def _build_citizen_authorization_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    principal = get("principal", "×××")
    principal_info = get("principal_info", "男/女，××××年××月××日出生，×族，……(写明工作单位和职务或者职业)，住……。联系方式：……")
    lawyer_agent = get("lawyer_agent", "×××，××律师事务所律师，联系方式：……")
    citizen_agent = get("citizen_agent", "×××，男/女，××××年××月××日出生，×族，……(写明工作单位和职务或者职业)，住……。联系方式：……。受托人系委托人的……(写明受托人与委托人的关系)")
    case_summary = get("case_summary", "……(写明当事人和案由)")
    agent_names = get("agent_names", "×××、×××")
    agent_1_name = get("agent_1_name", "×××")
    agent_1_authority = _normalize_multiline_items(fields.get("agent_1_authority"))
    agent_2_name = get("agent_2_name", "×××")
    agent_2_authority = _normalize_multiline_items(fields.get("agent_2_authority"))
    date = get("date", "××××年××月××日")

    return "\n".join([
        "授权委托书",
        "",
        f"委托人：{principal}，{principal_info}。",
        f"受委托人：{lawyer_agent}。",
        f"受委托人： {citizen_agent}。",
        f"现委托{agent_names}在{case_summary}一案中，作为我方参加诉讼的委托诉讼代理人。",
        "委托事项与权限如下：",
        f"委托诉讼代理人{agent_1_name}的代理事项和权限：",
        agent_1_authority,
        f"委托诉讼代理人{agent_2_name}的代理事项和权限：",
        agent_2_authority,
        "",
        "委托人(签名)    ",
        date,
        "",
        "【说明】",
        "1．本样式根据《中华人民共和国民事诉讼法》第四十九条、第五十八条、第五十九条以及《最高人民法院关于适用〈中华人民共和国民事诉讼法〉的解释》第七十八条、第八十五条制定，供公民当事人、法定代理人、共同诉讼代表人委托诉讼代理人参加诉讼用。",
        "2．当事人、法定代理人、共同诉讼代表人可以委托一至二人作为诉讼代理人。当事人有权委托诉讼代理人，提出回避申请，收集、提供证据，进行辩论，请求调解，提出上诉，申请执行。",
        "3．下列人员可以被委托为诉讼代理人：(一)律师、基层法律服务工作者；(二)当事人的近亲属或者工作人员；(三)当事人所在社区、单位及有关社会团体推荐的公民。",
        "4．与当事人有夫妻、直系血亲、三代以内旁系血亲、近姻亲关系以及其他有抚养、赡养关系的亲属，可以当事人近亲属的名义作为诉讼代理人。",
        "5．诉讼代理人除根据民事诉讼法第五十九条规定提交授权委托书外，还应当按照下列规定向人民法院提交相关材料：(一)律师应当提交律师执业证、律师事务所证明材料；(二)基层法律服务工作者应当提交法律服务工作者执业证、基层法律服务所出具的介绍信以及当事人一方位于本辖区内的证明材料；(三)当事人的近亲属应当提交身份证件和与委托人有近亲属关系的证明材料；(四)当事人的工作人员应当提交身份证件和与当事人有合法劳动人事关系的证明材料；(五)当事人所在社区、单位推荐的公民应当提交身份证件、推荐材料和当事人属于该社区、单位的证明材料；(六)有关社会团体推荐的公民应当提交身份证件和符合本解释第八十七条规定条件的证明材料。",
        "6．授权委托书必须记明委托事项和权限。诉讼代理人代为承认、放弃、变更诉讼请求，进行和解，提起反诉或者上诉，必须有委托人的特别授权。",
    ])


def _build_contract_body(fields):
    fields = fields or {}

    def get(key, default="……"):
        value = _clean_complaint_field(fields.get(key))
        return value or default

    contract_type = get("contract_type", "买卖合同")
    party_a = get("party_a", "甲方")
    party_b = get("party_b", "乙方")
    date = get("date", "××××年××月××日")

    if contract_type == "房屋租赁合同":
        clauses = [
            ("一、房屋基本情况", get("house", "房屋坐落、面积、用途及附属设施详见双方确认的交接清单。")),
            ("二、租赁期限", get("term", "租赁期限由双方约定。")),
            ("三、租金、押金及支付方式", get("rent", "租金、押金及支付方式由双方约定。")),
            ("四、费用承担", get("fees", "租赁期间相关费用由双方按照约定承担。")),
            ("五、房屋使用、维修和返还", "乙方应按照约定用途合理使用房屋；租赁期满或合同解除时，乙方应按约返还房屋及附属设施。"),
            ("六、违约责任", get("breach", "任何一方违反本合同约定的，应向守约方承担继续履行、采取补救措施或者赔偿损失等违约责任。")),
            ("七、争议解决", get("dispute", "因本合同发生争议，双方应协商解决；协商不成的，依法向有管辖权的人民法院起诉。")),
        ]
    elif contract_type == "借款合同":
        clauses = [
            ("一、借款金额", get("amount", "借款金额由双方约定。")),
            ("二、借款用途", get("purpose", "借款用途由双方约定。")),
            ("三、借款期限", get("term", "借款期限由双方约定。")),
            ("四、利息及还款方式", get("repayment", "利息、还款日期和还款方式由双方约定。")),
            ("五、担保", get("guarantee", "担保方式由双方约定；无担保的，应明确为无担保。")),
            ("六、违约责任", get("breach", "乙方未按约还款的，应承担逾期利息、违约金及甲方实现债权的合理费用。")),
            ("七、争议解决", "因本合同发生争议，双方应协商解决；协商不成的，依法向有管辖权的人民法院起诉。"),
        ]
    elif contract_type == "服务合同":
        clauses = [
            ("一、服务内容", get("content", "乙方应按照本合同约定向甲方提供服务。")),
            ("二、服务期限", get("term", "服务期限由双方约定。")),
            ("三、服务费用及支付方式", get("fee", "服务费用、支付节点和发票事项由双方约定。")),
            ("四、交付与验收", get("acceptance", "乙方完成服务成果后，甲方应按照约定标准进行验收。")),
            ("五、保密义务", "双方对在履行本合同过程中知悉的商业秘密、个人信息及其他非公开信息负有保密义务。"),
            ("六、违约责任", get("breach", "任何一方违反本合同约定的，应承担继续履行、采取补救措施或者赔偿损失等违约责任。")),
            ("七、争议解决", get("dispute", "因本合同发生争议，双方应协商解决；协商不成的，依法向有管辖权的人民法院起诉。")),
        ]
    elif contract_type == "劳动合同":
        clauses = [
            ("一、合同期限", get("term", "劳动合同期限由双方依法约定。")),
            ("二、工作内容和工作地点", get("position", "乙方岗位、工作内容和工作地点由双方约定。")),
            ("三、工作时间和休息休假", get("hours", "甲方依法安排乙方工作时间和休息休假。")),
            ("四、劳动报酬", get("salary", "甲方应按照约定和法律规定向乙方支付劳动报酬。")),
            ("五、社会保险和福利待遇", get("benefits", "甲方依法为乙方缴纳社会保险，乙方依法享受相关福利待遇。")),
            ("六、劳动保护、劳动条件和职业危害防护", "甲方应为乙方提供符合国家规定的劳动安全卫生条件和必要的劳动防护用品。"),
            ("七、解除、终止和违约责任", get("breach", "双方解除、终止劳动合同及违约责任按照法律法规和本合同约定执行。")),
        ]
    else:
        contract_type = "买卖合同"
        clauses = [
            ("一、标的物", get("subject", "标的物的名称、规格、数量、质量标准由双方约定。")),
            ("二、价款及支付方式", f"{get('price', '价款由双方约定')}。{get('payment', '付款方式由双方约定')}。"),
            ("三、交付与验收", get("delivery", "交付时间、地点、方式及验收标准由双方约定。")),
            ("四、质量保证", "甲方应保证交付的标的物符合合同约定、国家标准及通常使用目的。"),
            ("五、违约责任", get("breach", "任何一方违反本合同约定的，应承担继续履行、采取补救措施或者赔偿损失等违约责任。")),
            ("六、争议解决", get("dispute", "因本合同发生争议，双方应协商解决；协商不成的，依法向有管辖权的人民法院起诉。")),
        ]

    lines = [
        contract_type,
        "",
        f"甲方：{party_a}",
        f"乙方：{party_b}",
        "",
        "甲乙双方在平等、自愿、公平、诚实信用的基础上，经协商一致，订立本合同，共同遵守。",
        "",
    ]
    for title, content in clauses:
        lines.extend([title, content, ""])
    lines.extend([
        "其他",
        "本合同未尽事宜，双方可另行签订补充协议。补充协议与本合同具有同等法律效力。",
        "本合同自双方签字或盖章之日起生效。",
        "",
        "甲方(签字或盖章)：",
        "乙方(签字或盖章)：",
        "",
        date,
    ])
    return "\n".join(lines)


def _demo_generate(doc_type, description):
    if doc_type == "起诉状":
        fields = _extract_complaint_fields_from_description(description)
        body = _build_civil_complaint_body(fields)
        return {
            "demo_mode": True,
            "title": "民事起诉状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对法院、当事人身份信息、副本份数和证据材料。"],
            "document_type": doc_type,
        }
    if doc_type == "上诉状":
        fields = _extract_appeal_fields_from_description(description)
        body = _build_civil_appeal_body(fields)
        return {
            "demo_mode": True,
            "title": "民事上诉状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对原审法院、裁判文号、上诉期限、上诉请求和副本份数。"],
            "document_type": doc_type,
        }
    if doc_type == "答辩状":
        fields = _extract_defense_fields_from_description(description)
        body = _build_civil_defense_body(fields)
        return {
            "demo_mode": True,
            "title": "民事答辩状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对案号、答辩期限、答辩意见、证据材料和副本份数。"],
            "document_type": doc_type,
        }
    if doc_type == "申请执行书":
        fields = _extract_execution_fields_from_description(description)
        body = _build_execution_application_body(fields)
        return {
            "demo_mode": True,
            "title": "申请执行书",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对生效法律文书、执行内容、被执行人履行情况和管辖法院。"],
            "document_type": doc_type,
        }
    if doc_type == "反诉状":
        fields = _extract_counterclaim_fields_from_description(description)
        body = _build_civil_counterclaim_body(fields)
        return {
            "demo_mode": True,
            "title": "民事反诉状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对反诉请求、事实理由、证据材料和副本份数。"],
            "document_type": doc_type,
        }
    if doc_type == "管辖权异议书":
        fields = _extract_jurisdiction_objection_fields_from_description(description)
        body = _build_jurisdiction_objection_body(fields)
        return {
            "demo_mode": True,
            "title": "异议书",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对提出期限、原受理法院、案号、移送法院和管辖理由。"],
            "document_type": doc_type,
        }
    if doc_type == "司法确认申请书":
        fields = _extract_judicial_confirmation_fields_from_description(description)
        body = _build_judicial_confirmation_body(fields)
        return {
            "demo_mode": True,
            "title": "申请书",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对申请主体、调解协议内容、司法确认请求和管辖法院。"],
            "document_type": doc_type,
        }
    if doc_type == "公民授权委托书":
        fields = _extract_citizen_authorization_fields_from_description(description)
        body = _build_citizen_authorization_body(fields)
        return {
            "demo_mode": True,
            "title": "授权委托书",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对委托事项、代理权限和受托人身份材料。"],
            "document_type": doc_type,
        }
    if doc_type == "合同":
        fields = {"contract_type": "买卖合同"}
        body = _build_contract_body(fields)
        return {
            "demo_mode": True,
            "title": fields["contract_type"],
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对合同主体、主要条款、违约责任和争议解决条款。"],
            "document_type": doc_type,
        }

    return {
        "demo_mode": True,
        "title": f"{doc_type}",
        "header": {
            "court": "根据案件性质和标的额确定管辖法院",
            "原告": "甲方（根据案情填写）",
            "被告": "乙方（根据案情填写）",
        },
        "body": f"""民事{doc_type}

原告：[姓名/名称]，[身份证号/统一社会信用代码]，住所地：[地址]，联系电话：[电话]
被告：[姓名/名称]，[身份证号/统一社会信用代码]，住所地：[地址]，联系电话：[电话]

诉讼请求：
一、请求判令被告[具体请求内容]；
二、请求判令被告承担本案全部诉讼费用。

事实与理由：
（请根据以下案情描述完善事实部分）
{description}

综上所述，原告的合法权益受到侵害，为维护原告合法权益，特依法提起诉讼，恳请贵院支持原告全部诉讼请求。

此致
[管辖]人民法院

具状人：
年  月  日""",
        "attachments": ["身份证/营业执照复印件", "相关合同/协议", "支付凭证", "通讯记录/邮件", "其他证据材料"],
        "notes": [
            "上述为文书模板框架，请根据实际情况填写具体信息",
            "建议委托律师审查后再行提交",
            "注意诉讼时效，一般为3年",
            "管辖法院应根据被告住所地或合同履行地确定",
        ],
        "document_type": doc_type,
    }


def _demo_strategy(case_description):
    return {
        "demo_mode": True,
        "case_type": "合同纠纷",
        "cause_of_action": "买卖合同纠纷",
        "applicable_laws": [],
        "key_evidence": [
            "双方签订的合同原件或复印件",
            "交货凭证、验收单据",
            "付款凭证、银行流水",
            "双方往来邮件、微信聊天记录",
            "证人证言",
        ],
        "evidence_risks": ["部分电子证据需及时公证保全", "口头约定内容难以举证", "如需申请证人出庭需提前准备"],
        "legal_strategy": {
            "primary": "重点证明合同关系、己方履行情况、对方违约行为和实际损失，并根据完整材料确定请求。",
            "alternative": "如合同条款存在对己方不利的格式条款，可考虑主张该条款无效，以法律默认规则重新确定双方权利义务。",
            "settlement_advice": "在证据充分的情况下，可先行发送律师函催告履行；对方有履行意愿的，可协商分期付款方案；达成和解协议需确保可执行性。",
        },
        "jurisdiction_analysis": "合同纠纷一般由被告住所地或合同履行地人民法院管辖。如合同中有管辖约定，应审查该约定是否有效。",
        "statute_of_limitation": "诉讼时效的起算、中断和届满需结合完整事实与可核验法条单独判断。",
        "similar_cases": [],
        "success_probability": "需根据具体证据情况综合判断，建议在证据固定后评估（仅供参考）",
        "next_steps": [
            "收集并整理全部证据材料",
            "梳理案件时间线和关键事实",
            "计算各项损失的具体金额",
            "发送律师函催告对方履行",
            "如对方拒不履行，准备起诉材料",
        ],
    }


def _demo_student_legal(scenario, description):
    issue_label = scenario or "校园法律问题"
    return {
        "demo_mode": True,
        "issue_type": issue_label,
        "legal_relationship": "学生与学校之间通常同时存在教育管理关系、教育服务合同关系，以及人格权、财产权等民事权益保护关系；校外兼职还可能涉及劳动关系、劳务关系、居间服务或消费合同关系。",
        "school_rule_boundary": "学校可以依据学生手册、宿舍管理规定、奖助学金评定细则、社团管理办法等进行教育管理，但校规不得与法律法规相抵触，处理结果应当有事实依据、制度依据，并保障学生陈述、申辩和申诉权。",
        "applicable_laws": [],
        "rights_and_obligations": [
            "学生有权要求学校说明处理依据、事实认定、程序安排和救济渠道。",
            "学校处理奖助学金、宿舍调整、纪律处分、社团管理等事项时，应遵守公开、公平、公正和程序正当原则。",
            "涉及兼职被骗、押金不退、工资拖欠时，应尽快固定招聘信息、聊天记录、转账凭证、工作记录等证据。",
        ],
        "evidence_checklist": [
            "学生手册、学院通知、评定细则、处分或处理决定原文",
            "聊天记录、邮件、短信、会议通知、录音录像及其形成时间",
            "付款记录、押金收据、兼职招聘页面、考勤或工作成果",
            "证人姓名、同宿舍/同社团成员说明、辅导员或老师沟通记录",
        ],
        "risk_points": [
            {"point": "只口头沟通、没有留痕", "level": "中", "suggestion": "重要沟通尽量通过邮件、企业微信、短信等可留痕方式确认。"},
            {"point": "公开发帖点名指责", "level": "高", "suggestion": "避免使用侮辱性或未经核实的表述，防止引发名誉权纠纷。"},
            {"point": "错过校内申诉期限", "level": "高", "suggestion": "立即核对学生手册或处理决定载明的申诉期限。"},
        ],
        "action_plan": [
            "整理事实时间线，区分已经发生的事实、自己的诉求和可证明的证据。",
            "查阅并保存学校相关规章：学生手册、宿舍规定、奖助学金细则、社团管理办法等。",
            "先向辅导员、学院或相关职能部门提交书面沟通材料，要求说明依据和处理期限。",
            "校内处理不当时，按学生申诉委员会流程申诉；涉及劳动、诈骗、消费纠纷的，同步考虑劳动监察、市场监管、公安或法院途径。",
        ],
        "communication_template": f"老师/负责人您好：我是学生，就“{issue_label}”事项申请核实处理。事情经过如下：{description[:120]}。我希望学校说明处理依据、事实认定和可适用的申诉渠道，并请在合理期限内给予书面回复。相关证据我已整理，可按要求提交。谢谢。",
        "authority_channels": ["辅导员/学院学生工作办公室", "学校学生申诉处理委员会", "学生资助管理中心或奖助学金评审部门", "劳动监察、市场监管、公安机关或人民法院【视具体争议选择】"],
        "disclaimer": "以上为校园法律风险分析，具体校规条款和外部救济路径需结合学校正式文件、当地规定和完整证据进一步核实。",
    }


# ===== 辅助：LLM 调用 + 演示回退 =====
def _call_llm(system_prompt, user_prompt, temperature=0.2):
    """调用 LLM，自动使用当前用户绑定的模型"""
    return call_llm_json(system_prompt, user_prompt, temperature, _get_model_override(), _get_llm_config())


def _llm_or_demo(result, demo_func, *args):
    """如果 LLM 返回 demo_mode（API key 无效），自动回退到演示数据"""
    if not isinstance(result, dict):
        fallback = demo_func(*args)
        fallback["llm_error"] = "模型返回结构无效，已切换为安全演示结果"
        return fallback
    if result.get("demo_mode"):
        return demo_func(*args)
    if result.get("error"):
        fallback = demo_func(*args)
        fallback["llm_error"] = str(result.get("error", "模型调用失败"))[:200]
        return fallback
    return result


def _get_model_override():
    """获取当前用户绑定的模型"""
    if current_user.is_authenticated and current_user.llm_api_key and current_user.llm_model:
        return current_user.llm_model
    return ""


def _get_llm_config():
    """获取当前用户的个人 LLM 配置；未配置时回退全局配置。"""
    if current_user.is_authenticated:
        config = {}
        if current_user.llm_api_key:
            config["api_key"] = current_user.llm_api_key
        if current_user.llm_base_url:
            config["base_url"] = current_user.llm_base_url
        if current_user.llm_model:
            config["model"] = current_user.llm_model
        return config
    return {}


def _is_current_demo_mode():
    return is_demo_mode(_get_llm_config())


def _sse_stream(system_prompt, user_prompt, demo_func, *demo_args):
    """SSE 流式生成器：逐块发送文本，最后发送完整的 JSON 结果"""
    import time
    full_text = ""
    for item in call_llm_stream_json(system_prompt, user_prompt, 0.2, _get_model_override(), _get_llm_config()):
        if isinstance(item, str):
            full_text += item
            yield f"data: {json.dumps({'chunk': item})}\n\n"
        else:
            # 最终 dict
            if not isinstance(item, dict) and demo_func:
                item = demo_func(*demo_args)
                item["llm_error"] = "模型返回结构无效，已切换为安全演示结果"
            elif item.get("demo_mode") and demo_func:
                item = demo_func(*demo_args)
            elif item.get("error") and demo_func:
                item = demo_func(*demo_args)
                item["llm_error"] = item.pop("error", "")
            yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"


# ===== 明鉴法律智能体（统一入口） =====
def _legal_agent_request_data():
    if request.is_json:
        data = request.get_json(silent=True) or {}
    else:
        data = request.form.to_dict(flat=True)
        for key in ("document_fields", "client_messages"):
            value = data.get(key)
            if isinstance(value, str) and value.strip():
                try:
                    data[key] = json.loads(value)
                except (TypeError, ValueError):
                    data[key] = {} if key == "document_fields" else []
    return data, request.files.get("file")


def _agent_document_result(doc_type, fields):
    """Reuse the deterministic document builders without unverified template notes."""
    fields = dict(fields or {})
    if doc_type == "答辩状":
        title = "民事答辩状"
        body = _build_civil_defense_body(fields)
    elif doc_type == "上诉状":
        title = "民事上诉状"
        body = _build_civil_appeal_body(fields)
    elif doc_type == "合同":
        contract_type = _clean_complaint_field(fields.get("contract_type")) or "买卖合同"
        fields["contract_type"] = contract_type
        subject = _clean_complaint_field(fields.get("subject"))
        price = _clean_complaint_field(fields.get("price"))
        if "租赁" in contract_type:
            fields.setdefault("house", subject)
            fields.setdefault("rent", price)
        elif "借款" in contract_type:
            fields.setdefault("purpose", subject)
            fields.setdefault("amount", price)
        elif "服务" in contract_type:
            fields.setdefault("content", subject)
            fields.setdefault("fee", price)
        elif "劳动" in contract_type:
            fields.setdefault("position", subject)
            fields.setdefault("salary", price)
        title = contract_type
        body = _build_contract_body(fields)
    else:
        doc_type = "起诉状"
        title = "民事起诉状"
        body = _build_civil_complaint_body(fields)

    # Several legacy templates contain explanatory statute numbers.  The chat
    # agent never exports those until they have passed retrieval/audit.
    body = body.split("\n【说明】", 1)[0].rstrip()
    return {
        "title": title,
        "header": fields,
        "body": body,
        "attachments": [],
        "notes": [
            "这是根据你提供的信息形成的初稿，请核对主体、事实、请求、管辖和签署信息。",
            "未自动补写无法确认的法条、案例或程序期限。",
        ],
        "document_type": doc_type,
    }


def _agent_attachment_preview(filename, parsed):
    if parsed is None:
        return None
    session = RedactionSession()
    preview = session.redact(parsed.text or "")[:700]
    return {
        "filename": os.path.basename(filename or "附件"),
        "method": parsed.metadata.get("method", "unknown"),
        "page_count": parsed.metadata.get("page_count", 1),
        "char_count": parsed.metadata.get("char_count", len(parsed.text or "")),
        "ocr_confidence": parsed.metadata.get("ocr_confidence"),
        "warnings": list(parsed.metadata.get("warnings") or []),
        "preview": preview,
    }


def _agent_store_draft_state(conversation, doc_type, fields, status="collecting"):
    # Persist only field presence.  Raw personal data stays in the active
    # browser turn and must be re-sent for a final personalised export.
    safe_presence = {
        key: f"[已记录字段:{key}]"
        for key, value in dict(fields or {}).items()
        if str(value or "").strip()
    }
    state = load_state(conversation)
    state["draft"] = {
        "doc_type": doc_type,
        "fields": safe_presence,
        "status": status,
    }
    save_state(conversation, state)


def _agent_client_context(data, current_message):
    rows = data.get("client_messages") if isinstance(data, dict) else []
    values = []
    if isinstance(rows, list):
        for row in rows[-16:]:
            if not isinstance(row, dict) or row.get("role") != "user":
                continue
            content = row.get("content")
            if isinstance(content, str) and content.strip():
                values.append(content[:6000])
    if current_message and (not values or values[-1] != current_message):
        values.append(current_message)
    return "\n".join(values)[-24000:]


def _agent_needs_input(conversation, message, intent, answer, attachment=None, **extra):
    result = {
        "conversation_id": conversation.id,
        "intent": intent,
        "status": "needs_information",
        "answer": answer,
        "attachment": attachment,
        "suggested_follow_ups": [],
    }
    result.update(extra)
    add_message(conversation, "assistant", answer, result=result)
    save_record("legal_agent", message, result)
    return result


def _run_legal_agent_turn(data, upload=None):
    if not isinstance(data, dict):
        return {"error": "请求格式错误"}, 400
    message = data.get("message", "")
    conversation_id = data.get("conversation_id") or None
    requested_action = data.get("action", "")
    if not isinstance(message, str):
        return {"error": "消息内容必须是文本"}, 400
    if conversation_id is not None and not isinstance(conversation_id, str):
        return {"error": "conversation_id 必须是文本"}, 400
    if not isinstance(requested_action, str):
        return {"error": "action 必须是文本"}, 400

    parsed = None
    filename = ""
    if upload is not None and upload.filename:
        filename, parsed = parse_upload_detailed(upload)
    attachment_text = parsed.text if parsed is not None else ""
    if not message.strip() and not attachment_text.strip():
        return {"error": "请输入法律问题或上传需要分析的文件"}, 400
    if not message.strip() and attachment_text:
        message = "请分析这份附件，并告诉我关键问题和下一步建议。"

    try:
        conversation, _created = get_or_create_conversation(
            current_user.id, conversation_id, message
        )
    except ConversationNotFound as error:
        return {"error": str(error)}, 404

    state = load_state(conversation)
    active_draft = state.get("draft") if isinstance(state.get("draft"), dict) else {}
    active_doc_type = str(active_draft.get("doc_type") or "")
    attachment = _agent_attachment_preview(filename, parsed)
    add_message(conversation, "user", message, attachment=attachment)

    if active_doc_type and re.search(r"取消(?:文书|起草)?|退出(?:文书|起草)?|结束起草", message):
        state.pop("draft", None)
        save_state(conversation, state)
        answer = "已结束本轮文书起草。你可以继续咨询其他法律问题或重新发起一份文书。"
        result = {
            "conversation_id": conversation.id,
            "intent": INTENT_DRAFT,
            "status": "completed",
            "answer": answer,
            "suggested_follow_ups": ["咨询一个法律问题", "重新起草起诉状", "上传文件分析"],
        }
        assistant_message = add_message(conversation, "assistant", answer, result=result)
        result["message_id"] = assistant_message.id
        return result, 200

    intent = detect_intent(
        message,
        attachment_text=attachment_text,
        filename=filename,
        requested_action=requested_action,
        active_document_type=active_doc_type,
    )
    # Recommendation telemetry stores only a coarse legal domain, never the
    # consultation text, attachment content or extracted personal data.
    try:
        record_event(
            current_user.id,
            "consult",
            entity_type="conversation",
            entity_id=conversation.id,
            legal_domain=infer_legal_domain(message),
        )
    except Exception:
        db.session.rollback()
        app.logger.exception("Unable to record privacy-safe consultation interest")

    if intent == INTENT_DRAFT:
        doc_type = active_doc_type or infer_document_type(message) or "起诉状"
        previous_fields = active_draft.get("fields") if active_doc_type == doc_type else {}
        structured_fields = data.get("document_fields")
        if not isinstance(structured_fields, dict):
            structured_fields = {}
        client_context = _agent_client_context(data, message)
        fields = extract_document_fields(
            doc_type,
            client_context,
            previous_fields if isinstance(previous_fields, dict) else {},
            structured_fields,
        )
        missing = missing_document_fields(doc_type, fields)
        if missing:
            _agent_store_draft_state(conversation, doc_type, fields)
            answer = document_follow_up(doc_type, missing)
            return _agent_needs_input(
                conversation,
                message,
                intent,
                answer,
                attachment,
                document_type=doc_type,
                missing_fields=missing,
                collected_fields=list(fields),
            ), 200

        document = _agent_document_result(doc_type, fields)
        _agent_store_draft_state(conversation, doc_type, fields, status="complete")
        answer = (
            f"{document['title']}初稿已经生成。它只使用你提供的事实，没有自动补写未经核验的法条、案例或程序期限。"
            "请先核对当事人信息、请求、事实、管辖和证据，再导出使用。"
        )
        result = {
            "conversation_id": conversation.id,
            "intent": intent,
            "status": "completed",
            "answer": answer,
            "document": document,
            "can_export": True,
            "suggested_follow_ups": suggested_follow_ups(intent, document),
        }
        assistant_message = add_message(conversation, "assistant", answer, result=result)
        result["message_id"] = assistant_message.id
        save_record("legal_agent", json.dumps(fields, ensure_ascii=False), result)
        return result, 200

    combined_text = message
    if attachment_text:
        combined_text += f"\n\n【附件正文】\n{attachment_text}"
    minimum = 50 if intent == INTENT_REVIEW else 20 if intent == INTENT_ANALYZE else 1
    if len(combined_text.strip()) < minimum:
        label = "合同正文" if intent == INTENT_REVIEW else "文书正文"
        answer = f"请粘贴或上传完整的{label}后再分析；目前的信息不足，我不会猜测文件内容。"
        return _agent_needs_input(
            conversation, message, intent, answer, attachment
        ), 200

    max_chars = 10000 if intent == INTENT_REVIEW else 8000
    model_text, privacy_meta, document_meta = _prepare_document_for_model(
        combined_text, max_chars, parsed, filename
    )
    context = conversation_context(conversation, limit=20, char_budget=8000)
    verified_legal_anchors = legal_anchors_for_question(message)
    legal_basis_required = requires_legal_basis(message)
    try:
        if intent == INTENT_REVIEW:
            if _is_current_demo_mode():
                raw_result = _demo_review(model_text)
            else:
                prompt = REVIEW_CONTRACT_PROMPT.replace("{contract}", model_text)
                raw_result = _llm_or_demo(
                    _call_llm(SYSTEM_PROMPT, prompt), _demo_review, model_text
                )
            tool_result = raw_result
        elif intent == INTENT_ANALYZE:
            if _is_current_demo_mode():
                raw_result = _demo_analyze(model_text)
            else:
                prompt = ANALYZE_PROMPT.replace("{document}", model_text)
                raw_result = _llm_or_demo(
                    _call_llm(SYSTEM_PROMPT, prompt), _demo_analyze, model_text
                )
            tool_result = raw_result
        elif intent == INTENT_STRATEGY:
            if _is_current_demo_mode():
                raw_result = _demo_strategy(model_text)
            else:
                prompt = STRATEGY_PROMPT.replace("{case_description}", model_text)
                raw_result = _llm_or_demo(
                    _call_llm(SYSTEM_PROMPT, prompt), _demo_strategy, model_text
                )
            tool_result = raw_result
        else:
            focus = {
                INTENT_EVIDENCE: "证据清单整理",
                INTENT_RISK: "法律风险分析",
                INTENT_SEARCH: "法律法规检索",
            }.get(intent, "一般法律咨询")
            if _is_current_demo_mode():
                raw_result = _demo_search(model_text)
            else:
                memory_context = build_memory_context(current_user.id, conversation.id, model_text)
                prompt = build_general_prompt(
                    model_text, context, focus=focus,
                    conversation_summary=memory_context["summary"],
                    relevant_memories=memory_context["memories"],
                    attachment_evidence=attachment_text[:6000] if attachment_text else "无",
                    verified_legal_anchors=verified_legal_anchors,
                    legal_basis_required=legal_basis_required,
                )
                raw_result = _llm_or_demo(
                    _call_llm(SYSTEM_PROMPT, prompt), _demo_search, model_text
                )
            tool_result = raw_result
    except Exception as error:
        app.logger.exception("Legal agent capability failed")
        raw_result = _demo_search(model_text)
        raw_result["llm_error"] = str(error)[:200]
        tool_result = raw_result

    answer = compose_chat_answer(
        intent,
        tool_result,
        legal_basis_fallback=verified_legal_anchors,
        required_legal_basis=legal_basis_required,
    )
    result = {
        "conversation_id": conversation.id,
        "intent": intent,
        "status": "completed",
        "answer": answer,
        "attachment": attachment,
        "tool_result": tool_result,
        "suggested_follow_ups": suggested_follow_ups(intent, tool_result),
    }
    assistant_message = add_message(conversation, "assistant", answer, result=result)
    result["message_id"] = assistant_message.id
    save_record("legal_agent", model_text, result)
    return result, 200


@app.route("/api/legal-agent", methods=["POST"])
@require_approved
def legal_agent():
    """Unified synchronous legal-agent endpoint."""
    data, upload = _legal_agent_request_data()
    result, status = _run_legal_agent_turn(data, upload)
    return jsonify(result), status


@app.route("/api/legal-agent-stream", methods=["POST"])
@require_approved
def legal_agent_stream():
    """Unified SSE endpoint with progress, answer chunks and a final envelope."""
    data, upload = _legal_agent_request_data()
    if not isinstance(data, dict) or not isinstance(data.get("message", ""), str):
        return jsonify({"error": "消息内容必须是文本"}), 400

    def generate():
        yield f"data: {json.dumps({'stage': 'understand', 'status': '正在分析你的问题'}, ensure_ascii=False)}\n\n"
        yield f"data: {json.dumps({'stage': 'answer', 'status': '正在整理回答'}, ensure_ascii=False)}\n\n"
        result, status = _run_legal_agent_turn(data, upload)
        if status >= 400:
            yield f"data: {json.dumps({'done': True, 'result': result}, ensure_ascii=False)}\n\n"
            return
        answer = str(result.get("answer") or "")
        for index in range(0, len(answer), 56):
            yield f"data: {json.dumps({'chunk': answer[index:index + 56]}, ensure_ascii=False)}\n\n"
        yield f"data: {json.dumps({'done': True, 'result': result}, ensure_ascii=False)}\n\n"

    return Response(stream_with_context(generate()), mimetype="text/event-stream")


@app.route("/api/legal-agent/conversations", methods=["GET"])
@login_required
def legal_agent_conversations():
    return jsonify({"conversations": list_conversations(current_user.id)})


@app.route("/api/legal-agent/conversations/<conversation_id>", methods=["GET"])
@login_required
def legal_agent_conversation(conversation_id):
    try:
        conversation, _ = get_or_create_conversation(current_user.id, conversation_id)
    except ConversationNotFound as error:
        return jsonify({"error": str(error)}), 404
    return jsonify(serialize_conversation(conversation, include_messages=True))


@app.route("/api/legal-agent/conversations/<conversation_id>", methods=["DELETE"])
@login_required
def remove_legal_agent_conversation(conversation_id):
    if not delete_conversation(current_user.id, conversation_id):
        return jsonify({"error": "会话不存在或无权访问"}), 404
    return jsonify({"status": "ok"})


@app.route("/api/analyze", methods=["POST"])
@require_approved
def analyze_document():
    """分析法律文书：上传文件或粘贴文本"""
    parsed = None
    filename = ""
    if "file" in request.files and request.files["file"].filename:
        filename, parsed = parse_upload_detailed(request.files["file"])
        text = parsed.text
    else:
        data = request.get_json(silent=True) or {}
        text = request.form.get("text", "") or _text_field(data, "text")

    if text is None:
        return jsonify({"error": "文书内容必须是文本"}), 400
    if not text or len(text.strip()) < 20:
        return jsonify({"error": "请提供至少20字的文书内容"}), 400

    model_text, privacy_meta, document_meta = _prepare_document_for_model(
        text, 8000, parsed, filename
    )

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_analyze(model_text), privacy_meta, document_meta, "analyze"
        )
        save_record("analyze", model_text, result)
        return jsonify(result)

    prompt = ANALYZE_PROMPT.replace("{document}", model_text)
    try:
        result = _llm_or_demo(_call_llm(SYSTEM_PROMPT, prompt), _demo_analyze, model_text)
    except Exception as e:
        result = _demo_analyze(model_text)
        result["llm_error"] = str(e)[:200]

    result = _attach_processing_meta(result, privacy_meta, document_meta, "analyze")
    save_record("analyze", model_text, result)
    return jsonify(result)


@app.route("/api/analyze-stream", methods=["POST"])
@require_approved
def analyze_document_stream():
    """分析法律文书（流式）"""
    parsed = None
    filename = ""
    if "file" in request.files and request.files["file"].filename:
        filename, parsed = parse_upload_detailed(request.files["file"])
        text = parsed.text
    else:
        data = request.get_json(silent=True) or {}
        text = request.form.get("text", "") or _text_field(data, "text")

    if text is None:
        return jsonify({"error": "文书内容必须是文本"}), 400
    if not text or len(text.strip()) < 20:
        return jsonify({"error": "请提供至少20字的文书内容"}), 400

    model_text, privacy_meta, document_meta = _prepare_document_for_model(
        text, 8000, parsed, filename
    )

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_analyze(model_text), privacy_meta, document_meta, "analyze"
            )
            save_record("analyze", model_text, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = ANALYZE_PROMPT.replace("{document}", model_text)

    def gen():
        full_text = ""
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                full_text += item
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_analyze(model_text)
                item = _attach_processing_meta(item, privacy_meta, document_meta, "analyze")
                save_record("analyze", model_text, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 模块2: 法律法规智能检索 =====
@app.route("/api/search-provision", methods=["POST"])
@require_approved
def search_provision():
    """检索相关法律法规"""
    data = request.get_json(silent=True) or {}
    question = _text_field(data, "question")

    if question is None:
        return jsonify({"error": "法律问题必须是文本"}), 400
    if not question:
        return jsonify({"error": "请输入法律问题"}), 400

    model_question, privacy_meta, document_meta = _prepare_document_for_model(
        question, 4000
    )

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_search(model_question), privacy_meta, document_meta, "search"
        )
        save_record("search", model_question, result)
        return jsonify(result)

    prompt = SEARCH_PROVISION_PROMPT.replace("{question}", model_question)
    try:
        result = _llm_or_demo(_call_llm(SYSTEM_PROMPT, prompt), _demo_search, model_question)
    except Exception as e:
        result = _demo_search(model_question)
        result["llm_error"] = str(e)[:200]

    result = _attach_processing_meta(result, privacy_meta, document_meta, "search")

    save_record("search", model_question, result)
    return jsonify(result)


@app.route("/api/search-provision-stream", methods=["POST"])
@require_approved
def search_provision_stream():
    """检索法律法规（流式）"""
    data = request.get_json(silent=True) or {}
    question = _text_field(data, "question")

    if question is None:
        return jsonify({"error": "法律问题必须是文本"}), 400
    if not question:
        return jsonify({"error": "请输入法律问题"}), 400

    model_question, privacy_meta, document_meta = _prepare_document_for_model(
        question, 4000
    )

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_search(model_question), privacy_meta, document_meta, "search"
            )
            save_record("search", model_question, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = SEARCH_PROVISION_PROMPT.replace("{question}", model_question)

    def gen():
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_search(model_question)
                item = _attach_processing_meta(item, privacy_meta, document_meta, "search")
                save_record("search", model_question, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 模块3: 合同风险智能审查 =====
@app.route("/api/review-contract", methods=["POST"])
@require_approved
def review_contract():
    """审查合同风险"""
    parsed = None
    filename = ""
    if "file" in request.files and request.files["file"].filename:
        filename, parsed = parse_upload_detailed(request.files["file"])
        text = parsed.text
    else:
        data = request.get_json(silent=True) or {}
        text = _text_field(data, "contract")

    if text is None:
        return jsonify({"error": "合同内容必须是文本"}), 400
    if not text or len(text.strip()) < 50:
        return jsonify({"error": "请提供至少50字的合同内容"}), 400

    model_text, privacy_meta, document_meta = _prepare_document_for_model(
        text, 10000, parsed, filename
    )

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_review(model_text), privacy_meta, document_meta, "review"
        )
        save_record("review", model_text, result)
        return jsonify(result)

    prompt = REVIEW_CONTRACT_PROMPT.replace("{contract}", model_text)
    try:
        llm_result = _call_llm(SYSTEM_PROMPT, prompt)
        result = _llm_or_demo(llm_result, _demo_review, model_text)
    except Exception as e:
        result = _demo_review(model_text)
        result["llm_error"] = str(e)[:200]

    result = _attach_processing_meta(result, privacy_meta, document_meta, "review")
    save_record("review", model_text, result)
    return jsonify(result)


@app.route("/api/review-contract-stream", methods=["POST"])
@require_approved
def review_contract_stream():
    """审查合同风险（流式）"""
    parsed = None
    filename = ""
    if "file" in request.files and request.files["file"].filename:
        filename, parsed = parse_upload_detailed(request.files["file"])
        text = parsed.text
    else:
        data = request.get_json(silent=True) or {}
        text = _text_field(data, "contract")

    if text is None:
        return jsonify({"error": "合同内容必须是文本"}), 400
    if not text or len(text.strip()) < 50:
        return jsonify({"error": "请提供至少50字的合同内容"}), 400

    model_text, privacy_meta, document_meta = _prepare_document_for_model(
        text, 10000, parsed, filename
    )

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_review(model_text), privacy_meta, document_meta, "review"
            )
            save_record("review", model_text, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = REVIEW_CONTRACT_PROMPT.replace("{contract}", model_text)

    def gen():
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_review(model_text)
                item = _attach_processing_meta(item, privacy_meta, document_meta, "review")
                save_record("review", model_text, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 模块4: 法律文书智能生成 =====
@app.route("/api/generate-document", methods=["POST"])
@require_approved
def generate_document():
    """生成法律文书"""
    data = request.get_json(silent=True) or {}
    doc_type = _text_field(data, "doc_type", "起诉状")
    description = _text_field(data, "description")
    requirements = _text_field(data, "requirements", "格式规范，内容完整")
    complaint_fields = data.get("complaint_fields") or {}
    appeal_fields = data.get("appeal_fields") or {}
    defense_fields = data.get("defense_fields") or {}
    execution_fields = data.get("execution_fields") or {}
    counterclaim_fields = data.get("counterclaim_fields") or {}
    jurisdiction_objection_fields = data.get("jurisdiction_objection_fields") or {}
    judicial_confirmation_fields = data.get("judicial_confirmation_fields") or {}
    citizen_authorization_fields = data.get("citizen_authorization_fields") or {}
    contract_fields = data.get("contract_fields") or {}
    contract_type = _text_field(data, "contract_type")

    if None in (doc_type, description, requirements, contract_type):
        return jsonify({"error": "文书类型、案情描述和生成要求必须是文本"}), 400

    if doc_type == "起诉状":
        fields = complaint_fields if isinstance(complaint_fields, dict) else _extract_complaint_fields_from_description(description)
        body = _build_civil_complaint_body(fields)
        result = {
            "title": "民事起诉状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对法院、当事人身份信息、副本份数和证据材料。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "上诉状":
        fields = appeal_fields if isinstance(appeal_fields, dict) else _extract_appeal_fields_from_description(description)
        body = _build_civil_appeal_body(fields)
        result = {
            "title": "民事上诉状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对原审法院、裁判文号、上诉期限、上诉请求和副本份数。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "答辩状":
        fields = defense_fields if isinstance(defense_fields, dict) else _extract_defense_fields_from_description(description)
        body = _build_civil_defense_body(fields)
        result = {
            "title": "民事答辩状",
            "header": fields,
            "body": body,
            "attachments": [],
            "notes": ["请按实际情况核对案号、答辩期限、答辩意见、证据材料和副本份数。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "申请执行书":
        fields = execution_fields if isinstance(execution_fields, dict) else _extract_execution_fields_from_description(description)
        result = {
            "title": "申请执行书",
            "header": fields,
            "body": _build_execution_application_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对生效法律文书、执行内容、被执行人履行情况和管辖法院。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "反诉状":
        fields = counterclaim_fields if isinstance(counterclaim_fields, dict) else _extract_counterclaim_fields_from_description(description)
        result = {
            "title": "民事反诉状",
            "header": fields,
            "body": _build_civil_counterclaim_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对反诉请求、事实理由、证据材料和副本份数。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "管辖权异议书":
        fields = jurisdiction_objection_fields if isinstance(jurisdiction_objection_fields, dict) else _extract_jurisdiction_objection_fields_from_description(description)
        result = {
            "title": "异议书",
            "header": fields,
            "body": _build_jurisdiction_objection_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对提出期限、原受理法院、案号、移送法院和管辖理由。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "司法确认申请书":
        fields = judicial_confirmation_fields if isinstance(judicial_confirmation_fields, dict) else _extract_judicial_confirmation_fields_from_description(description)
        result = {
            "title": "申请书",
            "header": fields,
            "body": _build_judicial_confirmation_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对申请主体、调解协议内容、司法确认请求和管辖法院。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "公民授权委托书":
        fields = citizen_authorization_fields if isinstance(citizen_authorization_fields, dict) else _extract_citizen_authorization_fields_from_description(description)
        result = {
            "title": "授权委托书",
            "header": fields,
            "body": _build_citizen_authorization_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对委托事项、代理权限和受托人身份材料。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if doc_type == "合同":
        fields = contract_fields if isinstance(contract_fields, dict) else {}
        if contract_type and not fields.get("contract_type"):
            fields["contract_type"] = contract_type
        title = _clean_complaint_field(fields.get("contract_type")) or "买卖合同"
        result = {
            "title": title,
            "header": fields,
            "body": _build_contract_body(fields),
            "attachments": [],
            "notes": ["请按实际情况核对合同主体、主要条款、违约责任和争议解决条款。"],
            "document_type": doc_type,
        }
        save_record("generate", json.dumps(fields, ensure_ascii=False), result)
        return jsonify(result)

    if not description:
        return jsonify({"error": "请输入案情描述"}), 400

    generation_session = RedactionSession()
    model_description = generation_session.redact(description)[:8000]
    model_requirements = generation_session.redact(str(requirements))[:2000]
    privacy_meta = public_redaction_summary(generation_session)
    privacy_meta["applied_before_model"] = True
    document_meta = {
        "method": "structured_text_input",
        "page_count": 1,
        "char_count": len(description) + len(str(requirements)),
        "analyzed_char_count": len(model_description) + len(model_requirements),
        "truncated": len(description) > 8000 or len(str(requirements)) > 2000,
        "warnings": [],
        "ocr_confidence": None,
    }
    if document_meta["truncated"]:
        document_meta["warnings"].append("生成材料较长，已按字段上限截取后再分析。")

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_generate(doc_type, model_description),
            privacy_meta,
            document_meta,
            "generate",
        )
        save_record("generate", model_description, result)
        return jsonify(result)

    prompt = (GENERATE_DOCUMENT_PROMPT
              .replace("{doc_type}", doc_type)
              .replace("{description}", model_description)
              .replace("{requirements}", model_requirements))

    try:
        result = _llm_or_demo(
            _call_llm(SYSTEM_PROMPT, prompt),
            _demo_generate,
            doc_type,
            model_description,
        )
    except Exception as e:
        result = _demo_generate(doc_type, model_description)
        result["llm_error"] = str(e)[:200]

    result["document_type"] = doc_type
    _attach_processing_meta(
        result, privacy_meta, document_meta, "generate"
    )

    save_record("generate", model_description, result)
    return jsonify(result)


@app.route("/api/generate-document-stream", methods=["POST"])
@require_approved
def generate_document_stream():
    """生成法律文书（流式）"""
    data = request.get_json(silent=True) or {}
    doc_type = _text_field(data, "doc_type", "起诉状")
    description = _text_field(data, "description")
    requirements = _text_field(data, "requirements", "格式规范，内容完整")
    complaint_fields = data.get("complaint_fields") or {}
    appeal_fields = data.get("appeal_fields") or {}
    defense_fields = data.get("defense_fields") or {}
    execution_fields = data.get("execution_fields") or {}
    counterclaim_fields = data.get("counterclaim_fields") or {}
    jurisdiction_objection_fields = data.get("jurisdiction_objection_fields") or {}
    judicial_confirmation_fields = data.get("judicial_confirmation_fields") or {}
    citizen_authorization_fields = data.get("citizen_authorization_fields") or {}
    contract_fields = data.get("contract_fields") or {}
    contract_type = _text_field(data, "contract_type")

    if None in (doc_type, description, requirements, contract_type):
        return jsonify({"error": "文书类型、案情描述和生成要求必须是文本"}), 400

    if doc_type == "起诉状":
        fields = complaint_fields if isinstance(complaint_fields, dict) else _extract_complaint_fields_from_description(description)

        def gen_complaint():
            result = {
                "title": "民事起诉状",
                "header": fields,
                "body": _build_civil_complaint_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对法院、当事人身份信息、副本份数和证据材料。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_complaint()), mimetype="text/event-stream")

    if doc_type == "上诉状":
        fields = appeal_fields if isinstance(appeal_fields, dict) else _extract_appeal_fields_from_description(description)

        def gen_appeal():
            result = {
                "title": "民事上诉状",
                "header": fields,
                "body": _build_civil_appeal_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对原审法院、裁判文号、上诉期限、上诉请求和副本份数。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_appeal()), mimetype="text/event-stream")

    if doc_type == "答辩状":
        fields = defense_fields if isinstance(defense_fields, dict) else _extract_defense_fields_from_description(description)

        def gen_defense():
            result = {
                "title": "民事答辩状",
                "header": fields,
                "body": _build_civil_defense_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对案号、答辩期限、答辩意见、证据材料和副本份数。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_defense()), mimetype="text/event-stream")

    if doc_type == "申请执行书":
        fields = execution_fields if isinstance(execution_fields, dict) else _extract_execution_fields_from_description(description)

        def gen_execution():
            result = {
                "title": "申请执行书",
                "header": fields,
                "body": _build_execution_application_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对生效法律文书、执行内容、被执行人履行情况和管辖法院。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_execution()), mimetype="text/event-stream")

    if doc_type == "反诉状":
        fields = counterclaim_fields if isinstance(counterclaim_fields, dict) else _extract_counterclaim_fields_from_description(description)

        def gen_counterclaim():
            result = {
                "title": "民事反诉状",
                "header": fields,
                "body": _build_civil_counterclaim_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对反诉请求、事实理由、证据材料和副本份数。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_counterclaim()), mimetype="text/event-stream")

    if doc_type == "管辖权异议书":
        fields = jurisdiction_objection_fields if isinstance(jurisdiction_objection_fields, dict) else _extract_jurisdiction_objection_fields_from_description(description)

        def gen_jurisdiction_objection():
            result = {
                "title": "异议书",
                "header": fields,
                "body": _build_jurisdiction_objection_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对提出期限、原受理法院、案号、移送法院和管辖理由。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_jurisdiction_objection()), mimetype="text/event-stream")

    if doc_type == "司法确认申请书":
        fields = judicial_confirmation_fields if isinstance(judicial_confirmation_fields, dict) else _extract_judicial_confirmation_fields_from_description(description)

        def gen_judicial_confirmation():
            result = {
                "title": "申请书",
                "header": fields,
                "body": _build_judicial_confirmation_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对申请主体、调解协议内容、司法确认请求和管辖法院。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_judicial_confirmation()), mimetype="text/event-stream")

    if doc_type == "公民授权委托书":
        fields = citizen_authorization_fields if isinstance(citizen_authorization_fields, dict) else _extract_citizen_authorization_fields_from_description(description)

        def gen_citizen_authorization():
            result = {
                "title": "授权委托书",
                "header": fields,
                "body": _build_citizen_authorization_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对委托事项、代理权限和受托人身份材料。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_citizen_authorization()), mimetype="text/event-stream")

    if doc_type == "合同":
        fields = contract_fields if isinstance(contract_fields, dict) else {}
        if contract_type and not fields.get("contract_type"):
            fields["contract_type"] = contract_type

        def gen_contract():
            title = _clean_complaint_field(fields.get("contract_type")) or "买卖合同"
            result = {
                "title": title,
                "header": fields,
                "body": _build_contract_body(fields),
                "attachments": [],
                "notes": ["请按实际情况核对合同主体、主要条款、违约责任和争议解决条款。"],
                "document_type": doc_type,
            }
            save_record("generate", json.dumps(fields, ensure_ascii=False), result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen_contract()), mimetype="text/event-stream")

    if not description:
        return jsonify({"error": "请输入案情描述"}), 400

    generation_session = RedactionSession()
    model_description = generation_session.redact(description)[:8000]
    model_requirements = generation_session.redact(str(requirements))[:2000]
    privacy_meta = public_redaction_summary(generation_session)
    privacy_meta["applied_before_model"] = True
    document_meta = {
        "method": "structured_text_input",
        "page_count": 1,
        "char_count": len(description) + len(str(requirements)),
        "analyzed_char_count": len(model_description) + len(model_requirements),
        "truncated": len(description) > 8000 or len(str(requirements)) > 2000,
        "warnings": [],
        "ocr_confidence": None,
    }
    if document_meta["truncated"]:
        document_meta["warnings"].append("生成材料较长，已按字段上限截取后再分析。")

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_generate(doc_type, model_description),
                privacy_meta,
                document_meta,
                "generate",
            )
            save_record("generate", model_description, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = (GENERATE_DOCUMENT_PROMPT
              .replace("{doc_type}", doc_type)
              .replace("{description}", model_description)
              .replace("{requirements}", model_requirements))

    def gen():
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_generate(doc_type, model_description)
                item["document_type"] = doc_type
                _attach_processing_meta(
                    item, privacy_meta, document_meta, "generate"
                )
                save_record("generate", model_description, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 模块5: 案情策略分析 =====
@app.route("/api/strategy", methods=["POST"])
@require_approved
def case_strategy():
    """案情策略分析"""
    data = request.get_json(silent=True) or {}
    case_description = _text_field(data, "case_description")

    if case_description is None:
        return jsonify({"error": "案情描述必须是文本"}), 400
    if not case_description:
        return jsonify({"error": "请输入案情描述"}), 400

    model_description, privacy_meta, document_meta = _prepare_document_for_model(
        case_description, 8000
    )

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_strategy(model_description), privacy_meta, document_meta, "strategy"
        )
        save_record("strategy", model_description, result)
        return jsonify(result)

    prompt = STRATEGY_PROMPT.replace("{case_description}", model_description)

    try:
        result = _llm_or_demo(_call_llm(SYSTEM_PROMPT, prompt), _demo_strategy, model_description)
    except Exception as e:
        result = _demo_strategy(model_description)
        result["llm_error"] = str(e)[:200]

    result = _attach_processing_meta(result, privacy_meta, document_meta, "strategy")

    save_record("strategy", model_description, result)
    return jsonify(result)


@app.route("/api/strategy-stream", methods=["POST"])
@require_approved
def case_strategy_stream():
    """案情策略分析（流式）"""
    data = request.get_json(silent=True) or {}
    case_description = _text_field(data, "case_description")

    if case_description is None:
        return jsonify({"error": "案情描述必须是文本"}), 400
    if not case_description:
        return jsonify({"error": "请输入案情描述"}), 400

    model_description, privacy_meta, document_meta = _prepare_document_for_model(
        case_description, 8000
    )

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_strategy(model_description), privacy_meta, document_meta, "strategy"
            )
            save_record("strategy", model_description, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = STRATEGY_PROMPT.replace("{case_description}", model_description)

    def gen():
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_strategy(model_description)
                item = _attach_processing_meta(item, privacy_meta, document_meta, "strategy")
                save_record("strategy", model_description, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 模块6: 大学生法律问题咨询 =====
@app.route("/api/student-legal", methods=["POST"])
@require_approved
def student_legal():
    """大学生校园与校外兼职等法律问题咨询"""
    data = request.get_json(silent=True) or {}
    scenario = _text_field(data, "scenario", "校园管理")
    description = _text_field(data, "description")

    if scenario is None or description is None:
        return jsonify({"error": "问题类型和问题描述必须是文本"}), 400
    if not description or len(description) < 10:
        return jsonify({"error": "请至少输入10字的问题描述"}), 400

    model_description, privacy_meta, document_meta = _prepare_document_for_model(
        description, 8000
    )
    scenario_session = RedactionSession()
    model_scenario = scenario_session.redact(scenario)

    if _is_current_demo_mode():
        result = _attach_processing_meta(
            _demo_student_legal(model_scenario, model_description),
            privacy_meta,
            document_meta,
            "student_legal",
        )
        save_record("student_legal", model_description, result)
        return jsonify(result)

    prompt = (STUDENT_LEGAL_PROMPT
              .replace("{scenario}", model_scenario)
              .replace("{description}", model_description))

    try:
        result = _llm_or_demo(
            _call_llm(SYSTEM_PROMPT, prompt),
            _demo_student_legal,
            model_scenario,
            model_description,
        )
    except Exception as e:
        result = _demo_student_legal(model_scenario, model_description)
        result["llm_error"] = str(e)[:200]

    result = _attach_processing_meta(result, privacy_meta, document_meta, "student_legal")

    save_record("student_legal", model_description, result)
    return jsonify(result)


@app.route("/api/student-legal-stream", methods=["POST"])
@require_approved
def student_legal_stream():
    """大学生法律问题咨询（流式）"""
    data = request.get_json(silent=True) or {}
    scenario = _text_field(data, "scenario", "校园管理")
    description = _text_field(data, "description")

    if scenario is None or description is None:
        return jsonify({"error": "问题类型和问题描述必须是文本"}), 400
    if not description or len(description) < 10:
        return jsonify({"error": "请至少输入10字的问题描述"}), 400

    model_description, privacy_meta, document_meta = _prepare_document_for_model(
        description, 8000
    )
    scenario_session = RedactionSession()
    model_scenario = scenario_session.redact(scenario)

    if _is_current_demo_mode():
        def gen():
            result = _attach_processing_meta(
                _demo_student_legal(model_scenario, model_description),
                privacy_meta,
                document_meta,
                "student_legal",
            )
            save_record("student_legal", model_description, result)
            yield f"data: {json.dumps({'done': True, 'result': result})}\n\n"
        return Response(stream_with_context(gen()), mimetype="text/event-stream")

    prompt = (STUDENT_LEGAL_PROMPT
              .replace("{scenario}", model_scenario)
              .replace("{description}", model_description))

    def gen():
        for item in call_llm_stream_json(SYSTEM_PROMPT, prompt, 0.2, _get_model_override(), _get_llm_config()):
            if isinstance(item, str):
                yield f"data: {json.dumps({'chunk': item})}\n\n"
            else:
                if not isinstance(item, dict) or item.get("demo_mode") or item.get("error"):
                    item = _demo_student_legal(model_scenario, model_description)
                item = _attach_processing_meta(item, privacy_meta, document_meta, "student_legal")
                save_record("student_legal", model_description, item)
                yield f"data: {json.dumps({'done': True, 'result': item})}\n\n"
    return Response(stream_with_context(gen()), mimetype="text/event-stream")


# ===== 文档预处理与服务健康检查 =====
@app.route("/api/documents/extract", methods=["POST"])
@login_required
def extract_document():
    """Extract a file without consuming an AI call or persisting the source."""
    file_storage = request.files.get("file")
    if not file_storage or not file_storage.filename:
        return jsonify({"error": "请选择需要识别的文件"}), 400

    filename, parsed = parse_upload_detailed(file_storage)
    session = RedactionSession()
    redacted_text = session.redact(parsed.text)
    return jsonify({
        "filename": os.path.basename(filename),
        "extracted_text": parsed.text,
        "redacted_text": redacted_text,
        "document_meta": parsed.metadata,
        "privacy_meta": public_redaction_summary(session),
    })


@app.route("/api/privacy/preview", methods=["POST"])
@login_required
def privacy_preview():
    """Preview deterministic redaction; input and mapping are never stored."""
    data = request.get_json(silent=True) or {}
    source_text = _text_field(data, "text")
    if source_text is None:
        return jsonify({"error": "待脱敏内容必须是文本"}), 400
    if not source_text:
        return jsonify({"error": "请输入需要脱敏的文本"}), 400
    session = RedactionSession()
    return jsonify({
        "redacted_text": session.redact(source_text),
        "privacy_meta": public_redaction_summary(session),
    })


@app.route("/healthz", methods=["GET"])
def healthz():
    return jsonify({"status": "ok", "version": APP_VERSION})


@app.route("/readyz", methods=["GET"])
def readyz():
    components = {
        "database": False,
        "external_model": not is_demo_mode(),
        "ocr": ocr_available(),
    }
    try:
        db.session.execute(text("SELECT 1"))
        components["database"] = True
    except Exception:
        db.session.rollback()

    required_ready = components["database"] and components["external_model"]
    status = "ready" if required_ready and components["ocr"] else "degraded"
    return jsonify({
        "status": status,
        "version": APP_VERSION,
        "components": components,
    }), (200 if required_ready else 503)


# ===== API 状态检查 =====
@app.route("/api/status", methods=["GET"])
@login_required
def get_status():
    """检查 API 配置状态"""
    return jsonify({
        "demo_mode": _is_current_demo_mode(),
        "answer_source": "external_model_api",
        "ocr_available": ocr_available(),
        "privacy_redaction": True,
        "local_legal_knowledge_base": False,
        "version": APP_VERSION,
    })


@app.route("/api/history", methods=["GET"])
@login_required
def get_history():
    """获取当前用户最近的分析记录。"""
    def parse_result(value):
        try:
            return json.loads(value or "{}")
        except Exception:
            return {"raw": value or ""}

    records = (
        AnalysisRecord.query
        .filter_by(user_id=current_user.id)
        .order_by(AnalysisRecord.created_at.desc())
        .limit(20)
        .all()
    )
    safe_records = []
    for record in records:
        session = RedactionSession()
        safe_records.append({
            "id": record.id,
            "module_type": record.module_type,
            "input_text": session.redact(record.input_text or ""),
            "result": redact_nested_json(parse_result(record.result_json), session),
            "created_at": record.created_at.strftime("%Y-%m-%d %H:%M"),
        })
    return jsonify({"records": safe_records})


@app.route("/api/history/<int:record_id>", methods=["DELETE"])
@login_required
def delete_history(record_id):
    """删除当前用户自己的某条历史记录。"""
    record = AnalysisRecord.query.filter_by(id=record_id, user_id=current_user.id).first()
    if not record:
        return jsonify({"error": "历史记录不存在"}), 404
    db.session.delete(record)
    db.session.commit()
    return jsonify({"status": "ok"})


# ===== 获取文书模板信息 =====
@app.route("/api/document-templates", methods=["GET"])
@login_required
def get_templates():
    """获取文书模板列表"""
    with open("data/contract_risks.json", "r", encoding="utf-8") as f:
        data = json.load(f)
    return jsonify(data.get("document_templates", {}))


# ===== API 配置（运行时更新） =====
@app.route("/api/config", methods=["POST"])
@login_required
def set_config():
    """运行时更新 LLM 配置。管理员同步写全局配置，普通用户保存个人配置。"""
    data = request.get_json(silent=True) or {}
    api_key = _text_field(data, "api_key")
    base_url = _text_field(data, "base_url")
    model = _text_field(data, "model")
    if None in (api_key, base_url, model):
        return jsonify({"error": "API 配置字段必须是文本"}), 400
    if any("\n" in value or "\r" in value for value in (api_key, base_url, model)):
        return jsonify({"error": "API 配置字段不能包含换行符"}), 400
    if base_url and not base_url.startswith(("https://", "http://")):
        return jsonify({"error": "Base URL 必须使用 http:// 或 https://"}), 400

    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")

    def _update_env(key, value):
        if not value:
            return
        os.environ[key] = value
        # 读取现有 .env，更新或追加 key=value
        try:
            with open(env_path, "r", encoding="utf-8") as f:
                lines = f.readlines()
        except FileNotFoundError:
            lines = []
        found = False
        for i, line in enumerate(lines):
            if line.strip().startswith(f"{key}=") or line.strip().startswith(f"# {key}"):
                lines[i] = f"{key}={value}\n"
                found = True
                break
        if not found:
            lines.append(f"\n{key}={value}\n")
        with open(env_path, "w", encoding="utf-8") as f:
            f.writelines(lines)

    if api_key:
        current_user.llm_api_key = api_key
    if base_url:
        current_user.llm_base_url = base_url
    if model:
        current_user.llm_model = model
    db.session.commit()

    if current_user.is_admin and api_key:
        _update_env("LLM_API_KEY", api_key)
        update_config(api_key=api_key)
    if current_user.is_admin and base_url:
        _update_env("LLM_BASE_URL", base_url)
        update_config(base_url=base_url)
    if current_user.is_admin and model:
        _update_env("LLM_MODEL", model)
        update_config(model=model)

    return jsonify({"status": "ok", "message": "配置已保存"})


@app.route("/api/config", methods=["GET"])
@login_required
def get_config():
    """读取当前用户可见的 LLM 配置。API Key 只返回是否已设置。"""
    return jsonify({
        "base_url": current_user.llm_base_url or os.getenv("LLM_BASE_URL", "https://api.deepseek.com/v1"),
        "model": current_user.llm_model or os.getenv("LLM_MODEL", "deepseek-v4-pro"),
        "has_api_key": bool(current_user.llm_api_key),
    })


@app.route("/api/config/test", methods=["POST"])
@login_required
def test_config():
    """测试当前填写或已保存的 OpenAI 兼容配置。"""
    data = request.get_json(silent=True) or {}
    api_key_input = _text_field(data, "api_key")
    base_url_input = _text_field(data, "base_url")
    model_input = _text_field(data, "model")
    if None in (api_key_input, base_url_input, model_input):
        return jsonify({"error": "API 配置字段必须是文本"}), 400
    api_key = api_key_input or current_user.llm_api_key or ""
    base_url = base_url_input or current_user.llm_base_url or os.getenv("LLM_BASE_URL", "https://api.deepseek.com/v1")
    model = model_input or current_user.llm_model or os.getenv("LLM_MODEL", "deepseek-v4-pro")
    ok, message = test_llm_connection(api_key, base_url, model)
    return jsonify({"status": "ok" if ok else "error", "message": message}), (200 if ok else 400)


# ===== PDF 下载 =====

def _find_cjk_fonts():
    """查找系统中可用的中文字体，返回 (常规字体路径, 粗体/标题字体路径)"""
    import subprocess

    # 常见中文字体路径（按平台和优先级排列）
    search_paths = [
        # Windows
        ("C:/Windows/Fonts/simfang.ttf", "C:/Windows/Fonts/simhei.ttf"),
        ("C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"),
        ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc"),
        # Linux — wqy-microhei 优先（与 fpdf2 兼容性最好）
        ("/usr/share/fonts/wqy-microhei/wqy-microhei.ttc", "/usr/share/fonts/wqy-microhei/wqy-microhei.ttc"),
        ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
        ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"),
        ("/usr/share/fonts/truetype/arphic/uming.ttc", "/usr/share/fonts/truetype/arphic/ukai.ttc"),
        ("/usr/share/fonts/simsun.ttf", "/usr/share/fonts/simsun.ttf"),
        # macOS
        ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/PingFang.ttc"),
        ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ]

    for body_path, title_path in search_paths:
        if os.path.exists(body_path):
            return body_path, title_path if os.path.exists(title_path) else body_path

    # 通过 fc-list 查找（Linux）
    try:
        result = subprocess.run(
            ["fc-list", ":lang=zh", "file"],
            capture_output=True, text=True, timeout=5
        )
        if result.returncode == 0 and result.stdout.strip():
            fonts = []
            for line in result.stdout.strip().split("\n"):
                path = line.split(":")[0].strip()
                if path and os.path.exists(path):
                    fonts.append(path)
            if fonts:
                return fonts[0], fonts[0]
    except Exception:
        pass

    return None, None


def _format_party_info(info):
    """把结构化当事人信息转成文书中可读的一行。"""
    if not isinstance(info, dict):
        return str(info).strip()

    def pick(*keys):
        for key in keys:
            value = str(info.get(key, "")).strip()
            if value:
                return value
        return ""

    field_labels = [
        (pick("name", "姓名"), ""),
        (pick("gender", "sex", "性别"), ""),
        (pick("birthDate", "birth_date", "birthday", "出生日期"), "出生日期"),
        (pick("ethnicity", "nation", "民族"), ""),
        (pick("address", "住所", "住址", "地址"), "住"),
        (pick("idNumber", "id_number", "idCard", "身份证号", "身份证号码"), "身份证号"),
        (pick("phone", "mobile", "tel", "联系电话", "电话", "手机号"), "联系电话"),
    ]
    parts = []
    for value, label in field_labels:
        if not value:
            continue
        if label:
            if label == "出生日期":
                parts.append(f"{value}出生")
            else:
                parts.append(f"{label}：{value}" if label in ("身份证号", "联系电话") else f"{label}{value}")
        else:
            parts.append(value)

    return "，".join(parts).strip("，")


def _format_party_block(role, info):
    """把当事人信息拆成更适合法律文书 PDF 的短行。"""
    if not isinstance(info, dict):
        formatted = str(info).strip()
        return f"{role}：{formatted}" if formatted else f"{role}："

    def pick(*keys):
        for key in keys:
            value = str(info.get(key, "")).strip()
            if value:
                return value
        return ""

    name = pick("name", "姓名")
    gender = pick("gender", "sex", "性别")
    birth = pick("birthDate", "birth_date", "birthday", "出生日期")
    ethnicity = pick("ethnicity", "nation", "民族")
    address = pick("address", "住所", "住址", "地址")
    id_number = pick("idNumber", "id_number", "idCard", "身份证号", "身份证号码")
    phone = pick("phone", "mobile", "tel", "联系电话", "电话", "手机号")

    identity = "，".join(part for part in [name, gender, f"{birth}出生" if birth else "", ethnicity] if part)
    contact = "，".join(part for part in [
        f"住址：{address}" if address else "",
        f"身份证号：{id_number}" if id_number else "",
        f"联系电话：{phone}" if phone else "",
    ] if part)

    lines = []
    lines.append(f"{role}：{identity}。" if identity else f"{role}：")
    if contact:
        lines.append(f"　　{contact}。")
    return "\n".join(lines)


def _display_party_role(role):
    """把模型可能返回的英文角色名规范成中文。"""
    role_text = str(role).strip()
    role_map = {
        "plaintiff": "原告",
        "claimant": "原告",
        "applicant": "申请人",
        "defendant": "被告",
        "respondent": "被告",
        "appellee": "被上诉人",
        "appellant": "上诉人",
    }
    return role_map.get(role_text.lower(), role_text)


def _looks_like_party_line(text, role):
    return text.startswith(f"{role}：") or text.startswith(f"{role}:")


def _is_party_info_line(text):
    return any(_looks_like_party_line(text, role) for role in ("原告", "被告", "申请人", "被申请人", "上诉人", "被上诉人"))


def _normalize_party_line(text, header):
    """正文中若混入 plaintiff/defendant 结构化行，转成中文行。"""
    text = str(text).strip()
    if not isinstance(header, dict):
        return text
    parties = header.get("parties", {})
    if not isinstance(parties, dict):
        return text

    lowered = text.lower()
    for role, info in parties.items():
        role_text = str(role).strip()
        if lowered.startswith(f"{role_text.lower()}：") or lowered.startswith(f"{role_text.lower()}:"):
            display_role = _display_party_role(role_text)
            return _format_party_block(display_role, info)
    return text


def _is_standalone_numbering(text):
    """识别被模型单独换行输出的编号，如 1. / 1、/ 一、。"""
    if len(text) > 6:
        return False
    normalized = text.replace("．", ".").replace("。", ".")
    if normalized[:-1].isdigit() and normalized[-1] in ".、":
        return True
    return len(text) in (2, 3) and text[-1] in "、." and text[0] in "一二三四五六七八九十"


def _normalize_doc_line(text):
    """规范法律文书 PDF 行，避免常见生成格式破坏版式。"""
    s = str(text).strip().strip("\u3000").strip()
    if not s:
        return ""

    s = re.sub(r"^(\d+)[\.．。]\s*", r"\1、", s)
    s = re.sub(r"^([一二三四五六七八九十]+)[\.．。]\s*", r"\1、", s)
    s = re.sub(r"(?<=\d)%", "％", s)

    if s.startswith("日期：") or s.startswith("日期:"):
        s = s.split("：", 1)[-1] if "：" in s else s.split(":", 1)[-1]
        s = s.strip()

    if (s.startswith("原告：") or s.startswith("被告：")) and s[-1] not in "。；;":
        s += "。"

    return s


def _set_docx_run_font(run, size=12, bold=False):
    run.font.name = "仿宋"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    run.font.size = Pt(size)
    run.font.bold = bold


def _add_docx_paragraph(doc, text, align=None, first_line=False, bold=False, size=12, first_line_pt=24):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(first_line_pt)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    _set_docx_run_font(run, size=size, bold=bold)
    return paragraph


def _build_docx_from_body(title, body):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)

    style = doc.styles["Normal"]
    style.font.name = "仿宋"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    style.font.size = Pt(12)

    lines = str(body or "").splitlines()
    if lines and lines[0].strip() == title:
        lines = lines[1:]

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.font.name = "黑体"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_run.font.size = Pt(18)
    title_run.font.bold = True

    right_align_prefixes = ("起诉人", "上诉人(签名", "答辩人(签名", "申请执行人(签名", "反诉人(签名", "异议人(签名", "申请人(签名", "委托人(签名", "甲方(签字", "乙方(签字", "××××年", "    年", "年  月", "年    月")
    section_titles = ("诉讼请求：", "事实和理由：", "证据和证据来源，证人姓名和住所：", "上诉请求：", "上诉理由：", "【说明】")
    no_indent_prefixes = (
        "原告：", "法定代表人/主要负责人：", "委托诉讼代理人：", "被告：",
        "(以上写明", "此致", "附：", "1．", "2．", "3．", "4．", "5．", "6．"
    )

    previous_text = ""
    last_nonempty_text = ""
    for line in lines:
        text = line.strip()
        if not text:
            doc.add_paragraph()
            previous_text = ""
            continue
        if text == title:
            continue
        is_date_line = bool(re.match(r"^\d{4}年\d{1,2}月\d{1,2}日$", text))
        align = WD_ALIGN_PARAGRAPH.RIGHT if text.startswith(right_align_prefixes) or is_date_line else None
        first_line_pt = 24
        if title.endswith("合同"):
            is_section = bool(re.match(r"^[一二三四五六七八九十]+、", text)) or text == "其他"
            contract_no_indent_prefixes = ("甲方：", "乙方：", "甲乙双方", "本合同", "附：")
            first_line = not align and not is_section and not text.startswith(contract_no_indent_prefixes)
            first_line_pt = 24
        elif title in ("民事答辩状", "申请执行书", "民事反诉状", "异议书", "申请书", "授权委托书"):
            is_section = False
            first_line = not align and previous_text != "此致" and not (title == "申请书" and last_nonempty_text == "此致")
            first_line_pt = 28
        elif title == "民事上诉状":
            is_section = text in section_titles
            appeal_no_indent_prefixes = ("(以上写明", "此致", "附：", "1．", "2．", "3．", "4．", "5．", "6．")
            first_line = not align and not is_section and not text.startswith(appeal_no_indent_prefixes)
        else:
            is_section = text in section_titles
            first_line = not align and not is_section and not text.startswith(no_indent_prefixes)
        _add_docx_paragraph(doc, text, align=align, first_line=first_line, bold=is_section, first_line_pt=first_line_pt)
        previous_text = text
        last_nonempty_text = text

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@app.route("/api/download-docx", methods=["POST"])
@login_required
def download_docx():
    """将文书正文生成为 DOCX 文件下载"""
    data = request.get_json(silent=True) or {}
    title = data.get("title", "民事起诉状")
    body = data.get("body", "")
    if not body:
        doc_type = data.get("doc_type", "")
        if doc_type == "上诉状" or title == "民事上诉状":
            fields = data.get("appeal_fields") or {}
            body = _build_civil_appeal_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "答辩状" or title == "民事答辩状":
            fields = data.get("defense_fields") or {}
            body = _build_civil_defense_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "申请执行书" or title == "申请执行书":
            fields = data.get("execution_fields") or {}
            body = _build_execution_application_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "反诉状" or title == "民事反诉状":
            fields = data.get("counterclaim_fields") or {}
            body = _build_civil_counterclaim_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "管辖权异议书" or title == "异议书":
            fields = data.get("jurisdiction_objection_fields") or {}
            body = _build_jurisdiction_objection_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "司法确认申请书" or title == "申请书":
            fields = data.get("judicial_confirmation_fields") or {}
            body = _build_judicial_confirmation_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "公民授权委托书" or title == "授权委托书":
            fields = data.get("citizen_authorization_fields") or {}
            body = _build_citizen_authorization_body(fields if isinstance(fields, dict) else {})
        elif doc_type == "合同" or str(title).endswith("合同"):
            fields = data.get("contract_fields") or {}
            if isinstance(fields, dict) and not fields.get("contract_type"):
                fields["contract_type"] = data.get("contract_type") or title
            body = _build_contract_body(fields if isinstance(fields, dict) else {})
        else:
            fields = data.get("complaint_fields") or {}
            body = _build_civil_complaint_body(fields if isinstance(fields, dict) else {})
    buf = _build_docx_from_body(title, body)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"{title}.docx",
    )


@app.route("/api/download-pdf", methods=["POST"])
@login_required
def download_pdf():
    """将文书正文生成为 PDF 文件下载"""
    data = request.get_json(silent=True) or {}
    title = data.get("title", "法律文书")
    body = data.get("body", "")
    header = data.get("header", {})

    if not body:
        return jsonify({"error": "缺少文书正文"}), 400

    font_path, title_font = _find_cjk_fonts()

    if not font_path:
        return jsonify({"error": "服务器未安装中文字体，无法生成 PDF。请安装中文字体（如 fonts-wqy-zenhei）后重试。"}), 500

    # 预处理 body：合并断行的编号项，并清理模型输出中的伪空行/英文结构化行。
    lines = body.split("\n")
    merged = []
    for line in lines:
        s = _normalize_doc_line(_normalize_party_line(line, header))
        if s and merged and _is_standalone_numbering(merged[-1].strip()):
            # 上一行是编号（如 "1." "2."），合并到当前行
            merged[-1] = _normalize_doc_line(merged[-1].strip() + s)
        else:
            merged.append(s)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.set_margins(18, 22, 18)
    pdf.add_page()

    pdf.add_font("CJK", "", font_path)
    pdf.add_font("CJKBold", "", title_font)

    cell_w = pdf.w - pdf.l_margin - pdf.r_margin
    INDENT = "　　"  # 2 个全角空格 = 首行缩进
    BODY_FONT_SIZE = 11
    LINE_H = 7

    def write_cell(text, align="L", width=None):
        pdf.multi_cell(
            width or cell_w,
            LINE_H,
            str(text),
            align=align,
            new_x="LMARGIN",
            new_y="NEXT",
        )

    # 标题 — 居中
    pdf.set_font("CJKBold", "", 20)
    pdf.cell(cell_w, 14, title, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # 当事人信息。header 是结构化元数据；正文已包含正式抬头时不重复打印。
    if header:
        pdf.set_font("CJK", "", BODY_FONT_SIZE)
        printed_header = False
        court = header.get("court", "")
        has_closing_court = bool(court and (str(court) in body or "此致" in body))
        if court and not has_closing_court:
            write_cell(court)
            printed_header = True
        parties = header.get("parties", {})
        if isinstance(parties, dict):
            for role, info in parties.items():
                role_text = _display_party_role(role)
                if any(_looks_like_party_line(line.strip(), role_text) for line in merged):
                    continue
                write_cell(_format_party_block(role_text, info))
                printed_header = True
        if printed_header:
            pdf.ln(4)

    # 识别关键词
    section_markers = ("诉讼请求", "事实与理由", "理由", "答辩请求", "答辩意见",
                       "上诉请求", "上诉理由", "反诉请求", "证据清单", "附")
    closing_markers = ("此致", "谨呈")
    sig_markers = ("具状人", "起诉人", "答辩人", "上诉人", "代理人", "委托代理人",
                   "法定代理人")

    pdf.set_font("CJK", "", BODY_FONT_SIZE)

    for stripped in merged:
        if not stripped:
            pdf.ln(3)
            continue

        is_section = any(stripped.startswith(m) for m in section_markers)
        is_closing = any(stripped.startswith(m) for m in closing_markers)
        is_sig = any(stripped.startswith(s) for s in sig_markers)
        is_party = _is_party_info_line(stripped)
        is_date = stripped[:4].isdigit() and "年" in stripped[:8]

        if is_section:
            pdf.set_font("CJKBold", "", BODY_FONT_SIZE)
            write_cell(stripped)
            pdf.set_font("CJK", "", BODY_FONT_SIZE)
        elif is_party:
            write_cell(stripped)
        elif is_closing:
            pdf.ln(4)
            write_cell(stripped)
        elif is_sig or is_date:
            write_cell(stripped, align="R")
        elif stripped[0].isdigit() or (len(stripped) > 1 and stripped[0] in "一二三四五六七八九十"):
            # 编号项 — 不缩进
            write_cell(stripped)
        else:
            # 普通段落 — 首行缩进 2 全角字符
            write_cell(INDENT + stripped)

    # 补全落款（如果 body 没有此致/具状人）
    body_text = body
    if "此致" not in body_text and header:
        court = header.get("court", "")
        if court:
            pdf.ln(6)
            write_cell("此致")
            write_cell(court)

    if not any(s in body_text for s in sig_markers):
        pdf.ln(4)
        write_cell("具状人：", align="R")
        write_cell("      年    月    日", align="R")

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)

    filename = f"{title}.pdf"
    return send_file(
        buf,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )


# ===== 错误处理 =====
@app.errorhandler(404)
def not_found(e):
    return jsonify({"error": "路由不存在"}), 404


@app.errorhandler(500)
def server_error(e):
    app.logger.exception("Unhandled server error")
    return jsonify({"error": "服务器内部错误，请稍后重试"}), 500


if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        ensure_user_schema()
        admin_user = os.getenv("ADMIN_USERNAME", "admin")
        admin_pw = os.getenv("ADMIN_PASSWORD")
        u = User.query.filter_by(username=admin_user).first()
        if not u:
            generated_admin_pw = None
            if not admin_pw:
                generated_admin_pw = secrets.token_urlsafe(18)
                admin_pw = generated_admin_pw
            u = User(
                username=admin_user,
                password_hash=generate_password_hash(admin_pw),
                is_admin=True,
                is_approved=True,
            )
            db.session.add(u)
            db.session.commit()
            print(f"管理员账号已创建: {admin_user}")
            if generated_admin_pw:
                print(f"临时管理员密码: {generated_admin_pw}")
                print("请立即登录后修改密码，或在 .env 中设置 ADMIN_PASSWORD。")
        else:
            u.is_admin = True
            u.is_approved = True
            if admin_pw:
                u.password_hash = generate_password_hash(admin_pw)
            db.session.commit()
    print("=" * 60)
    print("  明鉴 - 基于大模型的法律文书智能助手")
    print("  访问地址: http://localhost:5000")
    print("=" * 60)
    socketio.run(app, debug=False, host="0.0.0.0", port=5000, allow_unsafe_werkzeug=True)
