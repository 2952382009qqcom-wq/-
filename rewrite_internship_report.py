from copy import deepcopy
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SRC = "认识实习报告_张硕_255150639.docx"
OUT = "认识实习报告_张硕_255150639_修改版.docx"


def set_paragraph_text(paragraph, text):
    """Replace text while keeping the first run's formatting as the local template."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    set_paragraph_text(paragraph, text)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text(paragraph, "")


def remove_outer_blank_columns(table):
    """Remove the empty helper columns that render as extra boxes on the cover form."""
    for row in table.rows:
        tr = row._tr
        cells = list(tr.tc_lst)
        if len(cells) >= 4:
            tr.remove(cells[-1])
            tr.remove(cells[0])

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is not None:
        grid_cols = list(tbl_grid.gridCol_lst)
        if len(grid_cols) >= 4:
            tbl_grid.remove(grid_cols[-1])
            tbl_grid.remove(grid_cols[0])


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.tcW
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def keep_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


doc = Document(SRC)

# Cover form: replace identity fields and remove redundant blank edge boxes.
cover = doc.tables[0]
set_cell_text(cover.cell(0, 2), "计2506班")
set_cell_text(cover.cell(1, 2), "张硕")
set_cell_text(cover.cell(2, 2), "255150639")
remove_outer_blank_columns(cover)
keep_table_width(cover, [1800, 3600])

paragraph_replacements = {
    13: (
        "本次认识实习围绕简易数字万用表的识别、焊接、装配和调试展开。实训中，我从最基础的元器件清点开始，逐步完成电阻、电容、二极管、三极管、蜂鸣器、显示屏、电池座、档位旋钮等部件的安装，并最终完成整机测试。整个过程既有电子工艺的规范要求，也有不少需要自己动手摸索的小细节。通过这次训练，我对万用表的结构、焊接工艺和电路调试有了更直观的认识，也真正体会到“图纸上看懂”和“手上做出来”不是一回事。"
    ),
    17: (
        "设备：恒温电烙铁、烙铁架、吸锡器、尖嘴镊子、斜口钳、十字螺丝刀、数字万用表、万用表套件电路板、防静电工作垫。"
    ),
    18: (
        "材料：电阻、电容、二极管、三极管、电位器、蜂鸣器、液晶显示屏、档位旋钮、电池扣、连接导线、焊锡丝、塑料外壳、固定螺丝及配套电池。"
    ),
    22: (
        "1. 前期准备：领取套件后先核对材料数量，对照电路板丝印和元件清单确认各元器件型号，避免一开始就把相近阻值或极性元件混在一起。"
    ),
    24: (
        "2. 元件识别：重点区分色环电阻阻值、电解电容正负极、二极管方向和三极管管脚位置，并把容易装错的元件单独放置。"
    ),
    26: (
        "3. 焊接安装：按照“先低后高、先小后大”的顺序进行焊接，先处理电阻和二极管，再安装电容、电位器、显示屏连接件等部位。"
    ),
    28: (
        "4. 焊点检查：每焊完一组元件就翻到板底查看焊点状态，重点检查虚焊、连锡、漏焊和引脚过长等问题，发现异常及时返修。"
    ),
    30: (
        "5. 整机装配：将主板、显示屏、旋钮、电池座和外壳逐步安装到位，装配时注意排线和导线不能被外壳挤压。"
    ),
    32: (
        "6. 通电测试：装入电池后依次测试直流电压、电阻和通断档，用已知电阻和干电池进行比对，判断读数是否稳定可靠。"
    ),
    36: (
        "焊接前先让电烙铁充分预热，并清理烙铁头氧化层，使焊锡能够顺利熔化。实际操作时，我采用“焊盘和引脚同时加热、焊锡随后送入”的方式，让焊锡自然铺展在焊盘上。专业上要求焊点应当呈小锥形或半圆形，表面有光泽、边缘过渡自然；说得直白一点，就是不能糊成一坨，也不能只沾一点点看着像没焊牢。"
    ),
    37: (
        "在焊接顺序上，我按照电路板标号逐项完成安装。电阻、电容等小元件比较密集，需要控制焊锡量；显示屏、电位器和电池座受力较多，焊点要更牢固。过程中我发现烙铁停留时间太短容易虚焊，时间太长又可能烫坏焊盘，所以需要把动作做稳，不急着赶进度。全部焊接结束后，我剪除多余引脚，并用肉眼和万用表通断档检查关键线路。"
    ),
    41: (
        "焊接检查通过后进入装配环节。首先把液晶显示屏放入外壳对应位置，确认显示窗口没有偏移；然后安装档位旋钮，使旋钮指示位置与内部触点对应；接着固定电池座和电路板，整理内部导线，最后合上外壳并拧紧螺丝。这个环节看起来比焊接简单，但实际上也很考验细心，因为外壳卡扣、旋钮方向和导线位置稍有不对，就会影响后续档位切换和开机效果。"
    ),
    45: (
        "整机完成后，我先检查外观是否平整、旋钮是否能顺畅转动，再装入电池进行通电测试。显示屏正常亮起后，分别用直流电压档测量干电池电压，用电阻档测量标准电阻，用通断档测试导线连通情况。测试结果总体稳定，读数与参考值偏差在可接受范围内，说明焊接和装配基本合格。通过测试我也明白了，作品做出来只是第一步，能稳定测量、能经得起检查才算真正完成。"
    ),
    55: (
        "这次认识实习让我对电子产品的制作流程有了比较完整的体验，不再只是停留在课堂上的电路符号和理论公式。"
    ),
    57: (
        "刚开始焊接时，我的手并不稳，焊锡量也控制不好，有时焊点太大，有时又担心没焊牢。后来在老师指导和反复练习中，我慢慢掌握了加热时间、送锡角度和元件固定的方法。专业一点说，这是对电子装配工艺规范的训练；用我自己的感受来说，就是越做越知道急不得，很多小问题都是因为图快造成的。"
    ),
    59: (
        "实训过程中，我也加深了对常见元器件的认识。比如电阻要看色环和阻值，二极管、电解电容必须注意极性，档位开关和显示屏安装还要考虑机械结构配合。以前觉得万用表只是一个测量工具，这次拆开、焊好、装起来之后，才发现它里面的每个小零件都有自己的作用。遇到显示异常或读数不稳时，我学会了从焊点、电源、元件方向和线路连通性几个角度逐步排查。"
    ),
    61: (
        "总的来说，本次实习既提升了我的动手能力，也让我意识到工程实践需要严谨、耐心和复盘。以后学习专业课程时，我会更重视理论和操作的结合，多动手验证，多记录问题，不把“差不多”当成标准。"
    ),
}

for idx, text in paragraph_replacements.items():
    set_paragraph_text(doc.paragraphs[idx], text)

problem_tables = [
    (
        "问题1描述",
        "焊接时个别焊点锡量偏多，靠得很近的焊盘之间出现连锡，导致通断测试时有短路提示。",
        "解决方法",
        "先停止继续安装，用吸锡器和烙铁清除多余焊锡，再重新补焊。后续焊接时减少单次送锡量，每完成几个焊点就检查一次，避免问题堆到最后才发现。",
    ),
    (
        "问题2描述",
        "色环电阻识别不够熟练，前期把两个阻值接近的电阻放错了位置，影响后续检测判断。",
        "解决方法",
        "重新对照色环表核对阻值，并使用万用表电阻档复测。安装前把元件按阻值分类摆放，做到“先确认、再插装、再焊接”。",
    ),
    (
        "问题3描述",
        "第一次通电时显示屏没有正常显示，检查发现电池座一处焊点不够饱满，供电接触不稳定。",
        "解决方法",
        "重新补焊电池座引脚，并检查正负极连接方向。补焊后再次装入电池，显示屏能够正常亮起，说明故障点已排除。",
    ),
    (
        "问题4描述",
        "外壳合上后档位旋钮转动略紧，个别档位切换时手感不顺，影响使用体验。",
        "解决方法",
        "拆开外壳重新调整旋钮与内部触点位置，清理板上残留的焊锡碎屑和剪脚残留物。重新装配后，旋钮切换明显顺畅。",
    ),
]

for table, rows in zip(doc.tables[1:], problem_tables):
    set_cell_text(table.cell(0, 0), rows[0])
    set_cell_text(table.cell(0, 1), rows[1])
    set_cell_text(table.cell(1, 0), rows[2])
    set_cell_text(table.cell(1, 1), rows[3])

doc.save(OUT)
print(OUT)
