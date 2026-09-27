"""Unit tests for in-memory document parsing and OCR fallbacks."""

from __future__ import annotations

import io
import types
import unittest
from unittest import mock

from core import doc_parser


class _Upload:
    def __init__(self, filename, content, content_length=None):
        self.filename = filename
        self._stream = io.BytesIO(content)
        self.content_length = content_length

    def read(self, size=-1):
        return self._stream.read(size)


class _Page:
    def __init__(self, text="", error=None):
        self.text = text
        self.error = error

    def extract_text(self):
        if self.error:
            raise self.error
        return self.text


def _reader_for(*texts):
    pages = [item if isinstance(item, _Page) else _Page(item) for item in texts]

    class Reader:
        is_encrypted = False

        def __init__(self, stream):
            self.stream = stream
            self.pages = pages

    return Reader


def _png_bytes(size=(2, 2)):
    from PIL import Image

    output = io.BytesIO()
    Image.new("RGB", size, "white").save(output, "PNG")
    return output.getvalue()


class MetadataTests(unittest.TestCase):
    def test_parsed_document_keeps_metadata_in_sync(self):
        result = doc_parser.ParsedDocument(
            text="合同文本", method="text", char_count=999, warnings=("提示",)
        )
        self.assertEqual(result.char_count, 4)
        self.assertEqual(
            result.metadata,
            {
                "method": "text",
                "page_count": 1,
                "char_count": 4,
                "warnings": ["提示"],
                "ocr_confidence": None,
            },
        )
        self.assertEqual(result.to_dict()["text"], "合同文本")

    def test_document_parse_error_exposes_http_status(self):
        error = doc_parser.DocumentParseError("不支持", 415, "unsupported")
        self.assertEqual(error.status_code, 415)
        self.assertEqual(error.http_status, 415)
        self.assertEqual(error.status, 415)
        self.assertEqual(
            error.to_dict(), {"error": "unsupported", "message": "不支持"}
        )


class TextAndUploadTests(unittest.TestCase):
    def test_txt_and_markdown_compatibility_api(self):
        self.assertEqual(doc_parser.parse_txt("你好".encode()), "你好")
        self.assertEqual(doc_parser.parse_file("README.MD", b"# title"), "# title")
        details = doc_parser.parse_file_detailed("note.txt", b"plain text")
        self.assertEqual(details.method, "text")
        self.assertEqual(details.char_count, 10)

    def test_gb18030_text_is_supported_with_warning(self):
        result = doc_parser.parse_file_detailed("note.txt", "合同".encode("gb18030"))
        self.assertEqual(result.text, "合同")
        self.assertTrue(result.warnings)

    def test_unknown_extension_and_missing_extension_are_415(self):
        for filename in ("payload.exe", "README"):
            with self.subTest(filename=filename):
                with self.assertRaises(doc_parser.DocumentParseError) as raised:
                    doc_parser.parse_file_detailed(filename, b"data")
                self.assertEqual(raised.exception.status_code, 415)
                self.assertEqual(raised.exception.code, "unsupported_media_type")

    def test_empty_and_oversized_documents_have_safe_statuses(self):
        with self.assertRaises(doc_parser.DocumentParseError) as empty:
            doc_parser.parse_file("empty.txt", b"")
        self.assertEqual(empty.exception.status_code, 422)

        with mock.patch.dict("os.environ", {"MAX_UPLOAD_MB": "0.000001"}):
            with self.assertRaises(doc_parser.DocumentParseError) as large:
                doc_parser.parse_file("large.txt", b"12")
        self.assertEqual(large.exception.status_code, 413)

    def test_upload_compatibility_and_detailed_api(self):
        upload = _Upload("memo.txt", b"abc")
        self.assertEqual(doc_parser.parse_upload(upload), ("memo.txt", "abc"))

        upload = _Upload(None, b"fallback")
        filename, details = doc_parser.parse_upload_detailed(upload)
        self.assertEqual(filename, "upload.txt")
        self.assertEqual(details.text, "fallback")

    def test_declared_upload_size_is_rejected_before_read(self):
        upload = mock.Mock(filename="big.txt", content_length=99)
        with mock.patch.dict("os.environ", {"MAX_UPLOAD_MB": "0.00001"}):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_upload(upload)
        self.assertEqual(raised.exception.status_code, 413)
        upload.read.assert_not_called()


class DocxTests(unittest.TestCase):
    def test_docx_includes_paragraphs_and_table_cells(self):
        paragraph = types.SimpleNamespace(text=" 条款一 ")
        row = types.SimpleNamespace(
            cells=[types.SimpleNamespace(text="甲方"), types.SimpleNamespace(text="乙方")]
        )
        table = types.SimpleNamespace(rows=[row])
        document = types.SimpleNamespace(paragraphs=[paragraph], tables=[table])
        fake_docx = types.SimpleNamespace(Document=mock.Mock(return_value=document))
        with mock.patch.dict("sys.modules", {"docx": fake_docx}):
            result = doc_parser.parse_file_detailed("contract.docx", b"docx")
            legacy_text = doc_parser.parse_docx(b"docx")
        self.assertEqual(result.text, "条款一\n甲方\t乙方")
        self.assertEqual(result.method, "docx")
        self.assertEqual(legacy_text, result.text)

    def test_docx_preserves_interleaved_paragraph_and_table_order(self):
        from docx import Document

        document = Document()
        document.add_paragraph("第一条")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲方"
        table.cell(0, 1).text = "乙方"
        document.add_paragraph("第二条")
        output = io.BytesIO()
        document.save(output)

        result = doc_parser.parse_file_detailed("ordered.docx", output.getvalue())

        self.assertEqual(result.text, "第一条\n甲方\t乙方\n第二条")

    def test_broken_docx_is_422(self):
        fake_docx = types.SimpleNamespace(Document=mock.Mock(side_effect=ValueError("bad zip")))
        with mock.patch.dict("sys.modules", {"docx": fake_docx}):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_file_detailed("broken.docx", b"broken")
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(raised.exception.code, "invalid_docx")


class PdfTests(unittest.TestCase):
    def test_digital_pdf_uses_text_layer_without_ocr(self):
        first_page = "第一页包含足够多的可提取合同正文，不需要启动光学识别。"
        second_page = "Second page contains enough searchable contract text."
        reader = _reader_for(first_page, second_page)
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.object(
            doc_parser, "_get_ocr_engine"
        ) as get_ocr:
            result = doc_parser.parse_file_detailed("case.pdf", b"%PDF")
        self.assertEqual(result.text, f"{first_page}\n{second_page}")
        self.assertEqual(result.method, "pdf_text")
        self.assertEqual(result.page_count, 2)
        self.assertIsNone(result.ocr_confidence)
        get_ocr.assert_not_called()

    def test_page_limit_is_413_before_extraction(self):
        first = mock.Mock()
        second = mock.Mock()
        reader = _reader_for(first, second)
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.dict(
            "os.environ", {"MAX_PDF_PAGES": "1"}
        ):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_pdf(b"%PDF")
        self.assertEqual(raised.exception.status_code, 413)

    def test_scanned_pdf_is_rendered_and_ocrd_in_memory(self):
        reader = _reader_for("")
        engine = mock.Mock(
            return_value=([[None, "扫描合同", 0.92], [None, "第二行", 0.88]], 0.01)
        )
        pdfium_document = mock.Mock()
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.object(
            doc_parser, "_get_ocr_engine", return_value=engine
        ), mock.patch.object(doc_parser, "_open_pdfium", return_value=pdfium_document), mock.patch.object(
            doc_parser, "_render_pdf_page", return_value=b"rendered-png"
        ) as render:
            result = doc_parser.parse_file_detailed("scan.pdf", b"%PDF")
        self.assertEqual(result.text, "扫描合同\n第二行")
        self.assertEqual(result.method, "pdf_ocr")
        self.assertEqual(result.ocr_confidence, 0.9)
        render.assert_called_once_with(pdfium_document, 0)

    def test_mixed_pdf_keeps_text_and_warns_when_ocr_is_unavailable(self):
        digital_text = "This digital page contains enough searchable contract text."
        reader = _reader_for(digital_text, "")
        missing = doc_parser.DocumentParseError(
            "OCR 功能不可用", 422, "ocr_unavailable"
        )
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.object(
            doc_parser, "_get_ocr_engine", side_effect=missing
        ):
            result = doc_parser.parse_file_detailed("mixed.pdf", b"%PDF")
        self.assertEqual(result.text, digital_text)
        self.assertEqual(result.method, "pdf_text")
        self.assertTrue(any("OCR 功能不可用" in item for item in result.warnings))

    def test_short_page_number_or_watermark_does_not_suppress_ocr(self):
        reader = _reader_for("1")
        engine = mock.Mock(return_value=([[None, "扫描页正文", 0.95]], 0.01))
        pdfium_document = mock.Mock()
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.object(
            doc_parser, "_get_ocr_engine", return_value=engine
        ), mock.patch.object(doc_parser, "_open_pdfium", return_value=pdfium_document), mock.patch.object(
            doc_parser, "_render_pdf_page", return_value=b"rendered-png"
        ):
            result = doc_parser.parse_file_detailed("watermark.pdf", b"%PDF")

        self.assertEqual(result.text, "扫描页正文")
        self.assertEqual(result.method, "pdf_ocr")

    def test_scanned_pdf_missing_ocr_is_a_clear_422(self):
        reader = _reader_for("")
        missing = doc_parser.DocumentParseError(
            "OCR 功能不可用", 422, "ocr_unavailable"
        )
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader), mock.patch.object(
            doc_parser, "_get_ocr_engine", side_effect=missing
        ):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_file_detailed("scan.pdf", b"%PDF")
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(raised.exception.code, "ocr_unavailable")

    def test_corrupt_pdf_is_422(self):
        reader = mock.Mock(side_effect=ValueError("broken xref"))
        with mock.patch.object(doc_parser, "_pdf_reader_class", return_value=reader):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_pdf(b"not a pdf")
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(raised.exception.code, "invalid_pdf")

    def test_pypdf_is_preferred_and_pypdf2_is_fallback(self):
        preferred = object()
        pypdf_module = types.SimpleNamespace(PdfReader=preferred)
        with mock.patch.object(
            doc_parser.importlib, "import_module", return_value=pypdf_module
        ) as importer:
            self.assertIs(doc_parser._pdf_reader_class(), preferred)
        importer.assert_called_once_with("pypdf")

        fallback = object()

        def import_module(name):
            if name == "pypdf":
                raise ImportError
            return types.SimpleNamespace(PdfReader=fallback)

        with mock.patch.object(doc_parser.importlib, "import_module", side_effect=import_module):
            self.assertIs(doc_parser._pdf_reader_class(), fallback)


class ImageAndOcrTests(unittest.TestCase):
    def tearDown(self):
        doc_parser._OCR_ENGINE = None

    def test_image_is_validated_then_ocrd(self):
        ocr_result = ("身份证明", 0.875)
        with mock.patch.object(doc_parser, "_run_ocr", return_value=ocr_result) as run:
            result = doc_parser.parse_file_detailed("photo.PNG", _png_bytes())
        self.assertEqual(result.text, "身份证明")
        self.assertEqual(result.method, "image_ocr")
        self.assertEqual(result.ocr_confidence, 0.875)
        run.assert_called_once()

    def test_pixel_limit_is_413_before_ocr(self):
        with mock.patch.dict("os.environ", {"MAX_IMAGE_PIXELS": "1"}), mock.patch.object(
            doc_parser, "_run_ocr"
        ) as run:
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_file_detailed("large.png", _png_bytes((2, 1)))
        self.assertEqual(raised.exception.status_code, 413)
        run.assert_not_called()

    def test_invalid_image_is_422(self):
        with self.assertRaises(doc_parser.DocumentParseError) as raised:
            doc_parser.parse_file_detailed("broken.jpg", b"not an image")
        self.assertEqual(raised.exception.status_code, 422)
        self.assertEqual(raised.exception.code, "invalid_image")

    def test_missing_rapidocr_is_not_an_internal_error(self):
        missing = doc_parser.DocumentParseError(
            "请安装 rapidocr-onnxruntime", 422, "ocr_unavailable"
        )
        with mock.patch.object(doc_parser, "_get_ocr_engine", side_effect=missing):
            with self.assertRaises(doc_parser.DocumentParseError) as raised:
                doc_parser.parse_file_detailed("photo.png", _png_bytes())
        self.assertEqual(raised.exception.status_code, 422)
        self.assertIn("rapidocr", str(raised.exception).lower())

    def test_ocr_engine_is_lazy_and_cached(self):
        engine = mock.Mock(return_value=([], 0.0))
        rapidocr = mock.Mock(return_value=engine)
        with mock.patch.object(doc_parser, "_rapidocr_class", return_value=rapidocr):
            self.assertIs(doc_parser._get_ocr_engine(), engine)
            self.assertIs(doc_parser._get_ocr_engine(), engine)
        rapidocr.assert_called_once_with()

    def test_ocr_available_does_not_load_models(self):
        with mock.patch.object(doc_parser.importlib.util, "find_spec", return_value=object()), mock.patch.object(
            doc_parser, "_get_ocr_engine"
        ) as get_engine:
            self.assertTrue(doc_parser.ocr_available())
        get_engine.assert_not_called()


if __name__ == "__main__":
    unittest.main()
