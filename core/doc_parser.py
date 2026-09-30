"""Safe in-memory document parsing with optional OCR support.

The compatibility API (``parse_txt``/``parse_pdf``/``parse_docx``/
``parse_file``/``parse_upload``) still returns plain text. New callers can use
the ``*_detailed`` variants to receive provenance and OCR metadata.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import importlib
import importlib.util
import io
import math
import os
from pathlib import Path
from statistics import mean
from typing import Any, Optional


DEFAULT_MAX_UPLOAD_MB = 15
DEFAULT_MAX_PDF_PAGES = 20
DEFAULT_MAX_IMAGE_PIXELS = 25_000_000
DEFAULT_MIN_PDF_TEXT_CHARS = 20
MAX_UPLOAD_FILES = 10

SUPPORTED_EXTENSIONS = frozenset(
    {"txt", "md", "pdf", "docx", "jpg", "jpeg", "png", "webp"}
)
IMAGE_EXTENSIONS = frozenset({"jpg", "jpeg", "png", "webp"})
_PDF_RENDER_SCALE = 2.0


class DocumentParseError(ValueError):
    """A document error that can be safely returned by an HTTP endpoint."""

    def __init__(
        self,
        message: str,
        status_code: int = 422,
        code: str = "document_parse_error",
    ) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code
        # Friendly aliases for integrations using another convention.
        self.http_status = status_code
        self.status = status_code
        self.code = code

    def to_dict(self) -> dict[str, Any]:
        return {"error": self.code, "message": self.message}


@dataclass
class ParsedDocument:
    """Text plus auditable details about how it was obtained."""

    text: str
    method: str
    page_count: int = 1
    char_count: int = 0
    warnings: list[str] = field(default_factory=list)
    ocr_confidence: Optional[float] = None
    files: list[dict[str, Any]] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.char_count = len(self.text)
        self.warnings = list(self.warnings)

    @property
    def metadata(self) -> dict[str, Any]:
        """Return JSON-ready metadata, excluding the potentially large text."""

        metadata = {
            "method": self.method,
            "page_count": self.page_count,
            "char_count": self.char_count,
            "warnings": list(self.warnings),
            "ocr_confidence": self.ocr_confidence,
        }
        if self.files:
            metadata.update(file_count=len(self.files), files=list(self.files))
        return metadata

    def to_dict(self) -> dict[str, Any]:
        return {"text": self.text, **self.metadata}


_OCR_ENGINE: Any = None


def _positive_number_env(name: str, default: float) -> float:
    raw = os.getenv(name)
    if raw is None:
        return default
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return default
    if not math.isfinite(value) or value <= 0:
        return default
    return value


def _positive_int_env(name: str, default: int) -> int:
    return max(1, int(_positive_number_env(name, float(default))))


def _max_upload_bytes() -> int:
    mb = _positive_number_env("MAX_UPLOAD_MB", float(DEFAULT_MAX_UPLOAD_MB))
    return max(1, int(mb * 1024 * 1024))


def _max_pdf_pages() -> int:
    return _positive_int_env("MAX_PDF_PAGES", DEFAULT_MAX_PDF_PAGES)


def _max_image_pixels() -> int:
    return _positive_int_env("MAX_IMAGE_PIXELS", DEFAULT_MAX_IMAGE_PIXELS)


def _min_pdf_text_chars() -> int:
    return _positive_int_env("MIN_PDF_TEXT_CHARS", DEFAULT_MIN_PDF_TEXT_CHARS)


def _as_bytes(content: bytes | bytearray | memoryview) -> bytes:
    if not isinstance(content, (bytes, bytearray, memoryview)):
        raise DocumentParseError(
            "上传内容必须是二进制数据。", 422, "invalid_document"
        )
    size = len(content)
    limit = _max_upload_bytes()
    if size > limit:
        raise DocumentParseError(
            f"文件过大（{size / 1024 / 1024:.1f} MB），最大允许 "
            f"{limit / 1024 / 1024:g} MB。",
            413,
            "file_too_large",
        )
    if size == 0:
        raise DocumentParseError("文件内容为空。", 422, "empty_document")
    return bytes(content)


def _extension(filename: str) -> str:
    safe_name = str(filename or "").replace("\\", "/")
    return Path(safe_name).suffix.lower().lstrip(".")


def _decode_text(content: bytes) -> tuple[str, list[str]]:
    encodings = ("utf-8-sig", "utf-8", "gb18030", "gbk", "gb2312")
    for index, encoding in enumerate(encodings):
        try:
            text = content.decode(encoding)
            warnings: list[str] = []
            if index > 1:
                warnings.append(f"文本使用 {encoding} 编码读取。")
            return text, warnings
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace"), [
        "无法确定文本编码，已替换无法解码的字符。"
    ]


def parse_txt(content: bytes) -> str:
    """Parse TXT/Markdown bytes and return text (compatibility API)."""

    data = _as_bytes(content)
    return _decode_text(data)[0]


def _pdf_reader_class():
    """Prefer the maintained pypdf package, with PyPDF2 compatibility."""

    try:
        return importlib.import_module("pypdf").PdfReader
    except (ImportError, AttributeError):
        try:
            return importlib.import_module("PyPDF2").PdfReader
        except (ImportError, AttributeError) as exc:
            raise DocumentParseError(
                "服务器缺少 PDF 解析组件（需要 pypdf 或 PyPDF2）。",
                422,
                "pdf_parser_unavailable",
            ) from exc


def _rapidocr_class():
    errors: list[BaseException] = []
    for module_name in ("rapidocr_onnxruntime", "rapidocr"):
        try:
            module = importlib.import_module(module_name)
            return module.RapidOCR
        except (ImportError, AttributeError) as exc:
            errors.append(exc)
    raise DocumentParseError(
        "OCR 功能不可用：请安装 rapidocr-onnxruntime（或 rapidocr）后重试。",
        422,
        "ocr_unavailable",
    ) from (errors[-1] if errors else None)


def ocr_available() -> bool:
    """Return whether a RapidOCR backend is importable without loading models."""

    for module_name in ("rapidocr_onnxruntime", "rapidocr"):
        try:
            if importlib.util.find_spec(module_name) is not None:
                return True
        except (ImportError, AttributeError, ValueError):
            try:
                module = importlib.import_module(module_name)
                if hasattr(module, "RapidOCR"):
                    return True
            except (ImportError, AttributeError):
                pass
    return False


def _get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        rapidocr = _rapidocr_class()
        try:
            _OCR_ENGINE = rapidocr()
        except Exception as exc:
            raise DocumentParseError(
                f"OCR 模型初始化失败：{exc}", 422, "ocr_unavailable"
            ) from exc
    return _OCR_ENGINE


def _normalise_ocr_output(raw: Any) -> tuple[str, Optional[float]]:
    """Normalise outputs from both RapidOCR Python package generations."""

    def score_values(values: Any) -> list[float]:
        if values is None:
            return []
        if isinstance(values, (str, bytes)):
            values = [values]
        else:
            try:
                values = list(values)
            except TypeError:
                values = [values]
        normalised: list[float] = []
        for value in values:
            try:
                score = float(value)
            except (TypeError, ValueError):
                continue
            if math.isfinite(score):
                normalised.append(score)
        return normalised

    if raw is None:
        return "", None

    texts_attr = getattr(raw, "txts", None)
    if texts_attr is not None:
        texts = [str(item).strip() for item in texts_attr if str(item).strip()]
        values = score_values(getattr(raw, "scores", None))
        return "\n".join(texts), mean(values) if values else None

    data = raw[0] if isinstance(raw, tuple) and len(raw) == 2 else raw
    if data is None:
        return "", None
    if isinstance(data, dict):
        texts_value = data.get("txts")
        if texts_value is None:
            texts_value = data.get("texts")
        if texts_value is None:
            texts_value = []
        texts = [str(item).strip() for item in texts_value if str(item).strip()]
        scores = score_values(data.get("scores"))
        return "\n".join(texts), mean(scores) if scores else None
    if not isinstance(data, (list, tuple)):
        return "", None

    texts: list[str] = []
    scores: list[float] = []
    for item in data:
        if isinstance(item, str):
            if item.strip():
                texts.append(item.strip())
            continue
        if isinstance(item, dict):
            text = item.get("text") or item.get("txt")
            score = item.get("score")
            if score is None:
                score = item.get("confidence")
        elif isinstance(item, (list, tuple)) and len(item) >= 2:
            text = item[-2]
            score = item[-1]
        else:
            continue
        if isinstance(text, str) and text.strip():
            texts.append(text.strip())
            scores.extend(score_values(score))
    return "\n".join(texts), mean(scores) if scores else None


def _run_ocr(image_data: bytes) -> tuple[str, Optional[float]]:
    try:
        raw = _get_ocr_engine()(image_data)
        text, confidence = _normalise_ocr_output(raw)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"OCR 识别失败：{exc}", 422, "ocr_failed"
        ) from exc
    if confidence is not None:
        confidence = round(max(0.0, min(1.0, confidence)), 4)
    return text.strip(), confidence


def _validate_image(content: bytes) -> tuple[int, int]:
    try:
        from PIL import Image
    except ImportError as exc:
        raise DocumentParseError(
            "图片解析功能不可用：服务器缺少 Pillow。",
            422,
            "image_parser_unavailable",
        ) from exc

    try:
        with Image.open(io.BytesIO(content)) as image:
            width, height = image.size
            detected_format = (image.format or "").upper()
            if detected_format not in {"JPEG", "PNG", "WEBP"}:
                raise DocumentParseError(
                    "图片内容不是有效的 JPG、PNG 或 WebP 文件。",
                    422,
                    "invalid_image",
                )
            pixel_count = width * height
            max_pixels = _max_image_pixels()
            if pixel_count > max_pixels:
                raise DocumentParseError(
                    f"图片像素过高（{pixel_count:,}），最大允许 {max_pixels:,} 像素。",
                    413,
                    "image_too_large",
                )
            image.verify()
            return width, height
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            "图片已损坏或格式不正确。", 422, "invalid_image"
        ) from exc


def _parse_image_detailed(content: bytes) -> ParsedDocument:
    _validate_image(content)
    text, confidence = _run_ocr(content)
    warnings = [] if text else ["OCR 未在图片中识别出文字。"]
    return ParsedDocument(
        text=text,
        method="image_ocr",
        page_count=1,
        warnings=warnings,
        ocr_confidence=confidence,
    )


def _close_quietly(value: Any) -> None:
    close = getattr(value, "close", None)
    if callable(close):
        try:
            close()
        except Exception:
            pass


def _render_pdf_page(pdf_document: Any, page_index: int) -> bytes:
    """Render one PDF page to PNG bytes entirely in memory."""

    page = bitmap = image = None
    try:
        page = pdf_document[page_index]
        width, height = page.get_size()
        pixel_count = math.ceil(width * _PDF_RENDER_SCALE) * math.ceil(
            height * _PDF_RENDER_SCALE
        )
        max_pixels = _max_image_pixels()
        if pixel_count > max_pixels:
            raise DocumentParseError(
                f"PDF 第 {page_index + 1} 页渲染像素过高（{pixel_count:,}），"
                f"最大允许 {max_pixels:,} 像素。",
                413,
                "image_too_large",
            )
        bitmap = page.render(scale=_PDF_RENDER_SCALE)
        image = bitmap.to_pil().convert("RGB")
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"PDF 第 {page_index + 1} 页无法渲染用于 OCR。",
            422,
            "pdf_render_failed",
        ) from exc
    finally:
        _close_quietly(image)
        _close_quietly(bitmap)
        _close_quietly(page)


def _open_pdfium(content: bytes):
    try:
        pdfium = importlib.import_module("pypdfium2")
    except ImportError as exc:
        raise DocumentParseError(
            "扫描版 PDF 需要 OCR，但服务器缺少 pypdfium2 渲染组件。",
            422,
            "pdf_ocr_unavailable",
        ) from exc
    try:
        return pdfium.PdfDocument(content)
    except Exception as exc:
        raise DocumentParseError(
            "PDF 已损坏，无法渲染扫描页。", 422, "invalid_pdf"
        ) from exc


def _parse_pdf_detailed(content: bytes) -> ParsedDocument:
    try:
        reader = _pdf_reader_class()(io.BytesIO(content))
        if getattr(reader, "is_encrypted", False):
            try:
                unlocked = reader.decrypt("")
            except Exception as exc:
                raise DocumentParseError(
                    "PDF 已加密，请先移除密码。", 422, "encrypted_pdf"
                ) from exc
            if not unlocked:
                raise DocumentParseError(
                    "PDF 已加密，请先移除密码。", 422, "encrypted_pdf"
                )
        page_count = len(reader.pages)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            "PDF 已损坏或格式不正确。", 422, "invalid_pdf"
        ) from exc

    max_pages = _max_pdf_pages()
    if page_count > max_pages:
        raise DocumentParseError(
            f"PDF 共 {page_count} 页，最大允许 {max_pages} 页。",
            413,
            "too_many_pages",
        )

    page_texts: list[str] = [""] * page_count
    scanned_pages: list[int] = []
    warnings: list[str] = []
    digital_pages = 0
    for index, page in enumerate(reader.pages):
        try:
            extracted = (page.extract_text() or "").strip()
        except Exception:
            extracted = ""
            warnings.append(f"第 {index + 1} 页文本层读取失败，已尝试 OCR。")
        if len("".join(extracted.split())) >= _min_pdf_text_chars():
            page_texts[index] = extracted
            digital_pages += 1
        else:
            page_texts[index] = extracted
            scanned_pages.append(index)

    confidences: list[float] = []
    ocr_pages = 0
    if scanned_pages:
        pdf_document = None
        try:
            # Initialise OCR before allocating rendered page bitmaps.
            _get_ocr_engine()
            pdf_document = _open_pdfium(content)
            for index in scanned_pages:
                image_data = _render_pdf_page(pdf_document, index)
                ocr_text, confidence = _run_ocr(image_data)
                if ocr_text:
                    page_texts[index] = ocr_text
                    ocr_pages += 1
                    if confidence is not None:
                        confidences.append(confidence)
                else:
                    warnings.append(f"第 {index + 1} 页 OCR 未识别出文字。")
        except DocumentParseError as exc:
            if exc.status_code == 413 or digital_pages == 0:
                raise
            warnings.append(f"部分扫描页未识别：{exc.message}")
        finally:
            _close_quietly(pdf_document)

    text = "\n".join(part for part in page_texts if part).strip()
    if ocr_pages and digital_pages:
        method = "pdf_text+ocr"
    elif ocr_pages:
        method = "pdf_ocr"
    else:
        method = "pdf_text"
    confidence = round(mean(confidences), 4) if confidences else None
    return ParsedDocument(
        text=text,
        method=method,
        page_count=page_count,
        warnings=warnings,
        ocr_confidence=confidence,
    )


def parse_pdf(content: bytes) -> str:
    """Extract text from a PDF, applying OCR only to scanned pages."""

    return _parse_pdf_detailed(_as_bytes(content)).text


def _parse_docx_detailed(content: bytes) -> ParsedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError(
            "服务器缺少 DOCX 解析组件（需要 python-docx）。",
            422,
            "docx_parser_unavailable",
        ) from exc

    try:
        document = Document(io.BytesIO(content))
        parts: list[str] = []
        body = getattr(getattr(document, "element", None), "body", None)
        if body is not None:
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            # Preserve the original paragraph/table order. Contract tables
            # often sit between clauses, so appending every table at the end
            # changes the meaning of the extracted document.
            for child in body.iterchildren():
                if child.tag.endswith("}p"):
                    value = Paragraph(child, document).text.strip()
                    if value:
                        parts.append(value)
                elif child.tag.endswith("}tbl"):
                    table = Table(child, document)
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append("\t".join(cells))
        else:
            # Lightweight compatibility path for alternate/fake Document
            # implementations used by integrations and unit tests.
            parts = [paragraph.text.strip() for paragraph in document.paragraphs]
            parts = [part for part in parts if part]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append("\t".join(cells))
    except Exception as exc:
        raise DocumentParseError(
            "DOCX 已损坏或格式不正确。", 422, "invalid_docx"
        ) from exc

    text = "\n".join(parts).strip()
    warnings = [] if text else ["DOCX 中未发现可提取的文字。"]
    return ParsedDocument(text=text, method="docx", page_count=1, warnings=warnings)


def parse_docx(content: bytes) -> str:
    """Extract paragraphs and table cells from a DOCX document."""

    return _parse_docx_detailed(_as_bytes(content)).text


def parse_file_detailed(filename: str, content: bytes) -> ParsedDocument:
    """Parse a supported in-memory file and return text with metadata."""

    extension = _extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        display = f".{extension}" if extension else "无扩展名"
        raise DocumentParseError(
            f"不支持的文件格式：{display}。支持 TXT、MD、PDF、DOCX、JPG、PNG、WebP。",
            415,
            "unsupported_media_type",
        )

    data = _as_bytes(content)
    if extension in {"txt", "md"}:
        text, warnings = _decode_text(data)
        return ParsedDocument(text=text, method="text", page_count=1, warnings=warnings)
    if extension == "pdf":
        return _parse_pdf_detailed(data)
    if extension == "docx":
        return _parse_docx_detailed(data)
    return _parse_image_detailed(data)


def parse_file(filename: str, content: bytes) -> str:
    """Parse a supported file and return plain text (compatibility API)."""

    return parse_file_detailed(filename, content).text


def _read_upload(file_storage: Any) -> bytes:
    limit = _max_upload_bytes()
    declared_length = getattr(file_storage, "content_length", None)
    if isinstance(declared_length, int) and declared_length > limit:
        raise DocumentParseError(
            f"文件过大，最大允许 {limit / 1024 / 1024:g} MB。",
            413,
            "file_too_large",
        )
    try:
        # One byte beyond the limit detects oversized untrusted streams while
        # preserving the no-temporary-file guarantee.
        content = file_storage.read(limit + 1)
    except TypeError:
        content = file_storage.read()
    except Exception as exc:
        raise DocumentParseError(
            "无法读取上传文件。", 422, "upload_read_failed"
        ) from exc
    return _as_bytes(content)


def parse_upload_detailed(file_storage: Any) -> tuple[str, ParsedDocument]:
    """Parse a Flask FileStorage-like object as ``(filename, details)``."""

    filename = getattr(file_storage, "filename", None) or "upload.txt"
    content = _read_upload(file_storage)
    return filename, parse_file_detailed(filename, content)


def parse_upload(file_storage: Any) -> tuple[str, str]:
    """Parse Flask FileStorage as ``(filename, text)`` (compatibility API)."""

    filename, result = parse_upload_detailed(file_storage)
    return filename, result.text


def parse_uploads_detailed(file_storages, *, parser=None) -> tuple[str, ParsedDocument]:
    """Parse ordered attachments atomically; preserve the single-file API."""
    uploads = [upload for upload in file_storages if getattr(upload, "filename", "")]
    if not uploads:
        raise DocumentParseError("请选择需要识别的文件。", 400, "missing_file")
    if len(uploads) > MAX_UPLOAD_FILES:
        raise DocumentParseError(f"一次最多上传 {MAX_UPLOAD_FILES} 个附件。", 400, "too_many_files")
    total_bytes = 0
    for upload in uploads:
        stream = getattr(upload, "stream", None)
        if stream is not None and stream.seekable():
            position = stream.tell()
            stream.seek(0, 2)
            total_bytes += stream.tell()
            stream.seek(position)
    if total_bytes > _max_upload_bytes():
        raise DocumentParseError(f"附件总大小不能超过 {_max_upload_bytes() / 1024 / 1024:g} MB。", 413, "files_too_large")

    parse = parser or parse_upload_detailed
    parts = []
    total_pages = 0
    for index, upload in enumerate(uploads, 1):
        try:
            filename, parsed = parse(upload)
        except DocumentParseError as error:
            raise DocumentParseError(
                f"第 {index} 个附件（{Path(upload.filename).name}）：{error.message}",
                error.status_code, error.code,
            ) from error
        parts.append((Path(filename).name, parsed))
        total_pages += parsed.page_count
        if total_pages > _max_pdf_pages():
            raise DocumentParseError("附件总页数超过允许的识别页数。", 413, "too_many_pages")
    if len(parts) == 1:
        return parts[0]
    warnings = [f"附件 {index}：{warning}" for index, (_, parsed) in enumerate(parts, 1) for warning in parsed.warnings]
    confidences = [parsed.ocr_confidence for _, parsed in parts if parsed.ocr_confidence is not None]
    combined = ParsedDocument(
        text="\n\n".join(f"【附件 {index}：{filename}】\n{parsed.text}" for index, (filename, parsed) in enumerate(parts, 1)),
        method="multi_file",
        page_count=total_pages,
        warnings=warnings,
        ocr_confidence=mean(confidences) if confidences else None,
        files=[{"filename": filename, **parsed.metadata} for filename, parsed in parts],
    )
    return f"{len(parts)} 个附件", combined


__all__ = [
    "DocumentParseError",
    "ParsedDocument",
    "SUPPORTED_EXTENSIONS",
    "ocr_available",
    "parse_docx",
    "parse_file",
    "parse_file_detailed",
    "parse_pdf",
    "parse_txt",
    "parse_upload",
    "parse_upload_detailed",
    "parse_uploads_detailed",
]
