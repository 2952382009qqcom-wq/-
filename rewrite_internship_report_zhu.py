from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SRC = "认识实习报告_张硕_255150639.docx"
OUT = "认识实习报告_朱俊琛_255150642_修改版.docx"


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    set_paragraph_text(paragraph, text)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text(paragraph, "")


def remove_outer_blank_columns(table):
    for row in table.rows:
        tr = row._tr
        cells = list(tr.tc_lst)
        if len(cells) >= 4:
            tr.remove(cells[-1])
            tr.remove(cells[0])

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is not None:
        cols = list(tbl_grid.gridCol_lst)
        if len(cols) >= 4:
            tbl_grid.remove(cols[-1])
            tbl_grid.remove(cols[0])


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.tcW
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def keep_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


doc = Document(SRC)

cover = doc.tables[0]
set_cell_text(cover.cell(0, 2), "计2506班")
set_cell_text(cover.cell(1, 2), "朱俊琛")
set_cell_text(cover.cell(2, 2), "255150642")
remove_outer_blank_columns(cover)
keep_table_width(cover, [1800, 3600])

paragraph_replacements = {
    13: (
        "本次认识实习的主要任务是完成数字万用表套件的焊接、安装和功能验证。实习从认识元器件开始，逐步涉及电路板识图、元件定位、焊接操作、外壳装配、量程切换检查以及基础测量测试。这个过程让我把平时听到的电阻、电容、二极管、显示屏、电池座等概念真正对应到实物上，也让我意识到电子产品不是简单把零件拼在一起，而是要在工艺、线路和结构之间保持一致。整体来看，这次实习既有专业训练，也很接地气，因为每一步都需要自己动手试、自己发现问题。"
    ),
    17: (
        "设备：恒温电烙铁、焊台、烙铁清洁海绵、吸锡器、数字万用表、镊子、斜口钳、小螺丝刀、元件盒和防静电工作台。"
    ),
    18: (
        "材料：万用表主板、色环电阻、瓷片电容、电解电容、二极管、三极管、电位器、液晶屏、蜂鸣器、档位盘、电池连接件、导线、焊锡丝、外壳及螺丝。"
    ),
    22: (
        "1. 材料核对：按照清单检查套件是否齐全，特别关注数量较多、外形相近的电阻和电容，防止后续安装时混用。"
    ),
    24: (
        "2. 图纸熟悉：先观察电路板丝印和元件位置，确认每个元件大致安装区域，再结合元器件极性要求进行分类摆放。"
    ),
    26: (
        "3. 分步焊接：从低矮元件开始焊接，依次完成电阻、二极管、电容、三极管、蜂鸣器和连接端子等部件。"
    ),
    28: (
        "4. 质量检查：焊接过程中随时查看焊点形态，对不饱满、发暗、连锡或疑似虚焊的位置及时返工处理。"
    ),
    30: (
        "5. 外壳安装：将显示屏、旋钮、主板和电池座装入外壳，确认按键和档位旋转位置没有偏斜。"
    ),
    32: (
        "6. 调试验收：通电后检查显示是否正常，再用电压档、电阻档和蜂鸣档分别测试，判断整机是否达到基本使用要求。"
    ),
    36: (
        "焊接操作是本次实习中最关键的一步。正式焊接前，我先让电烙铁达到稳定温度，并用湿海绵清理烙铁头，使其保持良好上锡状态。焊接时需要同时加热焊盘和元件引脚，再送入适量焊锡，让焊锡自然形成可靠连接。专业上讲，合格焊点应具备导电可靠、机械强度足够、外观圆润无毛刺等特点；用简单的话说，就是看着干净、焊得牢、不乱粘。"
    ),
    37: (
        "实际操作中，我明显感觉到焊接不是手一碰就能做好。烙铁角度偏了，焊锡就不容易流开；送锡多了，两个焊点可能连在一起；动作慢了，电路板又容易受热过久。因此我后面采用小范围、分批次的方式完成焊接，每完成一部分就停下来检查。这样虽然速度慢一些，但能减少返工，也更容易保证整体质量。"
    ),
    41: (
        "焊接完成后，我先检查主板背面的引脚是否修剪整齐，再进行整机装配。装配时要把液晶屏放正，确保屏幕窗口与外壳对齐；档位旋钮要和内部触点配合，否则转动时会出现错位；电池座和导线也要整理好，不能被外壳压住。这个步骤让我感受到，电子制作不仅看电路能不能通，还要考虑结构安装是否顺手、是否耐用。"
    ),
    45: (
        "测试阶段主要验证万用表能否正常显示和测量。装入电池后，显示屏可以正常启动，随后我用直流电压档测量干电池，用电阻档测量固定阻值电阻，并使用蜂鸣档检查导线连通。读数整体比较稳定，档位切换也能对应不同测量功能。虽然精度测试没有专业仪器那么严格，但从实训要求来看，成品已经能够完成基础测量任务。看到自己组装的万用表真正有读数时，还是挺有成就感的。"
    ),
    55: (
        "通过这次认识实习，我对电子工艺有了更直接的理解，也发现很多看似简单的操作其实都有规范要求。"
    ),
    57: (
        "刚开始我觉得万用表套件应该不难，按位置把元件焊上去就行。但真正操作后才发现，元件识别、焊点质量、极性方向、装配位置都会影响最终结果。比如一个电容方向装反，或者一个焊点虚焊，都可能导致整机无法正常工作。专业角度看，这是电子装配质量控制的问题；换成自己的话说，就是细节不过关，后面一定会找回来。"
    ),
    59: (
        "实习中我也学会了用排查思路解决问题。遇到显示不正常时，不能只想着是不是某个大部件坏了，而要从电源、焊点、元件方向、线路连通等方面一步步检查。这个过程锻炼了我的动手能力，也让我对课堂知识有了新的理解。以前学习电路时更多是看符号和公式，现在看到真实电路板后，感觉知识变得具体了很多。"
    ),
    61: (
        "总之，这次实习让我明白，工程实践需要认真、耐心和规范意识。以后在专业学习中，我会多关注实操训练，遇到问题多分析原因，不只满足于“做完”，而是尽量做到“做对、做好、能解释清楚”。"
    ),
}

for idx, text in paragraph_replacements.items():
    set_paragraph_text(doc.paragraphs[idx], text)

problem_tables = [
    (
        "问题1描述",
        "焊接初期烙铁头清理不及时，导致个别焊点表面发暗、不够光亮，看起来不够牢固。",
        "解决方法",
        "重新清洁烙铁头并补充少量焊锡，对发暗焊点进行复焊。之后每焊接一段时间就清理一次烙铁头，保证导热和上锡效果。",
    ),
    (
        "问题2描述",
        "安装电解电容时一开始没有特别注意正负极方向，差点把元件插反。",
        "解决方法",
        "暂停焊接后重新查看电容外壳标识和电路板丝印，确认极性后再安装。后续凡是有方向要求的元件，都先核对再焊接。",
    ),
    (
        "问题3描述",
        "第一次合上外壳后，显示屏位置略有偏移，窗口边缘遮住了一小部分显示区域。",
        "解决方法",
        "拆开外壳重新调整液晶屏卡槽位置，确认屏幕与外壳窗口平齐后再固定主板，重新装配后显示区域完整。",
    ),
    (
        "问题4描述",
        "测试电阻档时读数偶尔跳动，怀疑探针接触和焊点稳定性不够。",
        "解决方法",
        "先更换接触位置，再检查相关焊点并补焊疑似虚焊处。稳定探针接触后重新测量，读数波动明显减小。",
    ),
]

for table, rows in zip(doc.tables[1:], problem_tables):
    set_cell_text(table.cell(0, 0), rows[0])
    set_cell_text(table.cell(0, 1), rows[1])
    set_cell_text(table.cell(1, 0), rows[2])
    set_cell_text(table.cell(1, 1), rows[3])

doc.save(OUT)
print(OUT)
