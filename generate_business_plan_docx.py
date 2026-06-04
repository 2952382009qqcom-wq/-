# -*- coding: utf-8 -*-
"""生成明鉴项目商业计划书 Word 版。"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "明鉴_商业计划书_修改版.docx"

PRIMARY = RGBColor(13, 27, 42)
SECONDARY = RGBColor(27, 42, 74)
ACCENT = RGBColor(65, 90, 119)
BODY = RGBColor(51, 51, 51)
MUTED = RGBColor(120, 120, 120)
FILL = "F5F7FA"


def set_east_asia_font(run, name):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.tcW
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D6D6D6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9060):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.add_run("- ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    run_text = OxmlElement("w:r")
    run_text.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in [fld_begin, instr, fld_sep, run_text, fld_end]:
        run = paragraph.add_run()
        run._r.append(element)
    paragraph.add_run(" -")
    for run in paragraph.runs:
        set_east_asia_font(run, "宋体")
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Title", 22, PRIMARY),
        ("Heading 1", 16, PRIMARY),
        ("Heading 2", 13, SECONDARY),
        ("Heading 3", 12, ACCENT),
    ]:
        style = doc.styles[name]
        style.font.name = "黑体" if name != "Title" else "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(6 if name != "Heading 3" else 3)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header_p.add_run("明鉴 AI法律文书智能助手 - 商业计划书")
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_p)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(11)
    r.font.color.rgb = BODY
    r.bold = bold
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Pt(22)
    p.paragraph_format.first_line_indent = Pt(-11)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("● ")
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY
    r2 = p.add_run(text)
    set_east_asia_font(r2, "宋体")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = BODY


def add_numbered(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(22)
    p.paragraph_format.first_line_indent = Pt(-11)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{num}. ")
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY
    r.bold = True
    r2 = p.add_run(text)
    set_east_asia_font(r2, "宋体")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = BODY


def add_heading(doc, level, text):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], "0D1B2A")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_east_asia_font(r, "黑体")
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.bold = True
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            if row_i % 2 == 0:
                set_cell_shading(cells[j], FILL)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_east_asia_font(r, "宋体")
            r.font.size = Pt(9)
            r.font.color.rgb = BODY
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                set_cell_width(cell, width)
    doc.add_paragraph()
    return table


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("明  鉴")
    set_east_asia_font(r, "黑体")
    r.font.size = Pt(30)
    r.font.color.rgb = PRIMARY
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI法律文书智能助手")
    set_east_asia_font(r, "楷体")
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT

    info = [
        ("项目名称", "明鉴 — AI法律文书智能助手"),
        ("参赛赛道", "高教主赛道"),
        ("项目方向", "AI+法律科技（LegalTech）"),
        ("所属学校", "北京科技大学天津学院"),
        ("参赛赛事", "中国国际大学生创新大赛"),
        ("日期", "2026年5月"),
    ]
    doc.add_paragraph()
    table = add_table(doc, ["项目", "内容"], info, [2200, 6500])
    for row in table.rows[1:]:
        row.cells[0].paragraphs[0].runs[0].bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本文件为商业计划书，包含项目核心信息，未经授权不得对外传播")
    set_east_asia_font(r, "楷体")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, 1, "目  录")
    items = [
        "一、执行摘要",
        "二、项目背景与痛点分析",
        "三、产品 / 服务介绍",
        "四、市场分析",
        "五、商业模式",
        "六、营销策略",
        "七、团队介绍",
        "八、财务分析",
        "九、风险评估与对策",
        "十、发展规划",
        "十一、社会价值",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(item)
        set_east_asia_font(r, "黑体")
        r.font.size = Pt(13)
        r.font.color.rgb = PRIMARY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("注：整体目录结构保持原商业计划书不变")
    set_east_asia_font(r, "楷体")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    doc.add_page_break()


def build_doc():
    doc = Document()
    style_doc(doc)
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, 1, "一、执行摘要")
    add_heading(doc, 2, "1.1 项目概述")
    add_para(doc, "明鉴是一款面向普通公众、高校学生、基层法律服务场景和中小组织的 AI法律文书智能助手。项目以“法律依据可追溯、文书生成可落地、使用门槛足够低”为核心原则，提供文书分析、法规检索、合同审查、文书生成、案情策略分析和大学生法律问题咨询六大模块。")
    add_para(doc, "项目已经完成前后端主要功能开发：前端采用 HTML、CSS 和 JavaScript 构建响应式交互界面，提供模块导航、结构化表单、文件上传、流式输出、历史记录和学生专题表单；后端采用 Flask 构建 API 服务，接入用户认证、文件解析、本地知识库检索、模型调用、PDF 文书导出和数据记录等能力。")
    add_para(doc, "项目的政策方向与数字法治、公共法律服务体系建设和“人工智能+”行动相契合。明鉴不替代律师或学校正式处理程序，而是提供一个低门槛的前置辅助入口，帮助用户先完成事实梳理、依据查询、证据整理和文书初稿准备。")
    add_heading(doc, 2, "1.2 核心价值")
    add_bullet(doc, "产品价值：把复杂法律任务拆成可填写、可检索、可生成、可导出的操作流程，降低普通用户使用门槛。")
    add_bullet(doc, "工程价值：完成前端页面、后端接口、知识库检索、流式输出、PDF 导出、账号管理和历史记录等完整闭环。")
    add_bullet(doc, "场景价值：新增大学生法律问题模块，聚焦校园管理、宿舍纠纷、奖助学金、社团合规、校外兼职被骗和实习劳动争议。")
    add_heading(doc, 2, "1.3 项目现状")
    add_para(doc, "目前项目已完成六大核心模块并部署运行，具备从用户输入到结果展示的完整流程。学生法律模块已接入独立 Prompt、结构化表单、后端流式接口和专题报告展示，是当前版本最具辨识度的新增功能。")
    add_heading(doc, 2, "1.4 竞争优势")
    add_table(doc, ["维度", "明鉴", "传统法律数据库", "通用 AI助手"], [
        ["产品形态", "六大模块一体化", "偏资料检索", "需用户自行设计提示词"],
        ["前端体验", "结构化表单、文件上传、流式输出", "检索框为主", "聊天框为主"],
        ["后端能力", "Flask API、知识库、PDF 导出、历史记录", "数据库服务", "模型服务"],
        ["学生场景", "校园法律专题报告", "通常不覆盖", "需要用户自行组织问题"],
        ["定位", "普惠、教学、基层可用", "专业付费工具", "通用问答工具"],
    ], [1500, 2900, 2450, 2200])

    add_heading(doc, 1, "二、项目背景与痛点分析")
    add_heading(doc, 2, "2.1 社会背景")
    add_para(doc, "随着法治中国和数字中国建设持续推进，公共法律服务正在从线下窗口向线上平台、热线服务和智能辅助工具延伸。普通公众、高校学生和基层法律服务人员都存在基础法律信息整理、文书初稿生成和证据清单梳理的需求。")
    add_para(doc, "明鉴的切入点不是承诺替代专业法律服务，而是抓住高频、基础、可结构化的环节。用户可以先用系统整理事实、查询相关依据、生成沟通文本或文书初稿，再根据事项复杂程度选择校内申诉、行政投诉、律师咨询或司法途径。")
    add_heading(doc, 2, "2.2 三大核心痛点")
    add_heading(doc, 3, "痛点一：法律服务第一步门槛高")
    add_para(doc, "很多用户并不是一开始就需要完整代理服务，而是先想知道问题属于什么类型、需要准备哪些证据、应该找谁沟通以及能否形成一份规范材料。传统咨询成本较高，网络信息又容易碎片化。")
    add_heading(doc, 3, "痛点二：法律文书与校内材料不会写")
    add_para(doc, "起诉状、答辩状、投诉材料、情况说明和校内申诉文本都需要清楚表达事实、理由和诉求。普通用户和学生群体往往不知道哪些事实关键、哪些证据要列、表述如何更正式。")
    add_heading(doc, 3, "痛点三：学生场景缺少专门工具")
    add_para(doc, "大学生在宿舍、奖助学金、社团、兼职、实习和师生沟通中遇到的问题，往往同时涉及校规、合同、劳动、消费和人格权等因素。通用法律问答很少主动提醒校规边界、校内申诉路径和正式沟通文本。")
    add_heading(doc, 2, "2.3 现有解决方案的不足")
    add_para(doc, "传统法律数据库强调资料检索，面向律师和法务人员较多；通用 AI助手虽然使用方便，但缺少本地知识库约束和法律场景输出结构。明鉴希望在两者之间提供一个更适合普通用户、学生和基层场景的轻量工具。")

    add_heading(doc, 1, "三、产品 / 服务介绍")
    add_heading(doc, 2, "3.1 法律文书智能分析")
    add_para(doc, "用户上传或粘贴法律文书后，系统识别文书类型，提取当事人信息和关键条款，标注风险点，并给出修订方向。该模块适合合同、起诉状、律师函、判决书等材料的初步阅读。")
    add_heading(doc, 2, "3.2 法律法规智能检索")
    add_para(doc, "用户用自然语言提出问题，系统先在本地知识库中检索相关法条，再结合模型生成解释和实务建议。该模块保留关键词检索优化作为基础能力，但不把单一算法作为项目唯一卖点。")
    add_heading(doc, 2, "3.3 合同风险智能审查")
    add_para(doc, "用户上传合同后，系统依据合同风险模式库识别违约责任不对等、争议解决条款不利、单方解除权、缺失必要条款等常见问题，并输出修改建议。")
    add_heading(doc, 2, "3.4 法律文书智能生成")
    add_para(doc, "本模块采用结构化表单，用户按当事人信息、事实经过、核心诉求和额外要求填写内容，系统生成起诉状、答辩状、上诉状、律师函、合同和仲裁申请书，并支持 PDF 下载。")
    add_heading(doc, 2, "3.5 案情策略分析")
    add_para(doc, "用户填写当事人、事实、争议焦点和当前进展后，系统输出案件类型、适用依据、证据风险、诉讼策略、管辖分析和下一步行动建议。")
    add_heading(doc, 2, "3.6 大学生法律问题咨询")
    add_para(doc, "学生模块是项目最新重点。用户选择校园管理、宿舍纠纷、师生矛盾、奖学金/助学金、入党政审、社团合规、校外兼职被骗、实习劳动争议等问题类型后，按时间、地点、对象、事实、处理情况、证据和诉求填写信息。系统生成校园法律专题报告，内容包括法律关系、校规边界、证据清单、风险提醒、处理路径、沟通文本和可联系渠道。")
    add_para(doc, "该模块强调先校内沟通和申诉，再根据争议性质选择劳动监察、市场监管、公安或法院等外部途径。它既适合学生个人维权，也适合高校普法、辅导员初步分流和法学实训课程展示。")
    add_heading(doc, 2, "3.7 前后端与工程实现")
    add_para(doc, "前端部分完成侧边导航、六大模块页面、表单校验、文件上传、流式输出展示、结果卡片、历史记录面板和学生专题报告展示。学生模块使用多字段表单，引导用户把分散事实整理成可分析材料。")
    add_para(doc, "后端部分以 Flask 为核心，提供文书分析、法规检索、合同审查、文书生成、策略分析和学生法律咨询接口。系统支持本地知识库检索、风险模式匹配、模型兜底演示、用户认证、操作记录保存和 PDF 生成，具备可演示、可迭代的 MVP 基础。")

    add_heading(doc, 1, "四、市场分析")
    add_heading(doc, 2, "4.1 市场规模与趋势")
    add_para(doc, "本计划书采用政策与行业供给数据作为市场判断依据。司法部公开数据表明，截至 2024 年底，全国法律服务机构和专业法律服务人员规模持续扩大，公共法律服务实体平台每年办理大量业务。法律服务需求真实存在，并且正在向线上化、普惠化、智能化方向延伸。")
    add_para(doc, "在政策层面，国家持续推进数字法治、智慧司法和公共法律服务体系建设，人工智能与公共服务、民生服务场景结合的空间正在扩大。明鉴的定位与这一趋势一致。")
    add_heading(doc, 2, "4.2 目标用户画像")
    add_heading(doc, 3, "核心用户：普通公众")
    add_para(doc, "普通公众主要面对劳动纠纷、民间借贷、合同纠纷、婚姻家庭等高频问题，需求集中在法规查询、材料整理和文书初稿。")
    add_heading(doc, 3, "重点用户：高校学生与校园管理场景")
    add_para(doc, "高校学生是本项目重点强调的场景。学生常见问题包括宿舍矛盾、奖助学金争议、社团活动风险、校外兼职被骗、实习劳动争议和个人信息保护。学生模块能够把事实、证据、校规边界、沟通文本和救济路径整合为一份专题报告。")
    add_heading(doc, 3, "延伸用户：基层法律服务需求者")
    add_para(doc, "基层司法所、调解组织、法律援助窗口和高校法学实训场景需要便捷的法规检索、文书参考和风险提示工具。明鉴可作为辅助工具补充传统公共法律服务。")
    add_heading(doc, 2, "4.3 竞品分析")
    add_table(doc, ["产品", "主要能力", "学生场景", "价格/门槛"], [
        ["明鉴", "六大模块、前后端闭环、PDF 导出", "重点覆盖", "免费/低成本"],
        ["传统法律数据库", "法律法规和案例检索", "覆盖较弱", "付费门槛较高"],
        ["通用 AI助手", "通用问答与文本生成", "依赖用户提示", "免费/付费"],
        ["模板网站", "固定文书模板", "缺少分析", "低门槛"],
    ], [1900, 3500, 2300, 1700])
    add_heading(doc, 2, "4.4 市场机会总结")
    add_para(doc, "市场机会不在于夸大行业规模，而在于抓住高频、基础、可结构化的法律服务入口。明鉴通过学生模块切入高校真实场景，再延展到普通公众和基层法律服务，有利于形成清晰的赛事展示重点和后续落地路径。")

    add_heading(doc, 1, "五、商业模式")
    add_heading(doc, 2, "5.1 收入模式")
    add_heading(doc, 3, "基础版")
    add_para(doc, "基础版面向普通公众和高校学生，提供文书分析、基础法规检索、合同风险初筛和学生法律问题基础分析，用于积累用户和验证场景。")
    add_heading(doc, 3, "专业版")
    add_para(doc, "专业版面向高频法律需求者、基层法律工作者、中小企业主和高校实训团队，提供不限次检索、文书生成、深度策略分析、学生专题报告和优先支持。")
    add_heading(doc, 3, "机构版")
    add_para(doc, "机构版面向律所、企业法务、高校和基层法律服务机构，提供私有化部署、专属知识库、API 集成、模板定制和试点项目支持。")
    add_heading(doc, 2, "5.2 成本结构")
    add_bullet(doc, "研发成本：前端交互优化、后端接口维护、学生模块迭代、知识库更新。")
    add_bullet(doc, "算力成本：模型 API 调用、云服务器、文件解析和 PDF 生成。")
    add_bullet(doc, "合规成本：法律内容审核、隐私保护、免责声明和安全审计。")
    add_heading(doc, 2, "5.3 盈利预测")
    add_para(doc, "项目前期以真实使用反馈和场景验证为主要目标。第一阶段重点完成高校学生法律场景试用、法学实训合作和少量机构试点，以覆盖基础运行成本和验证付费意愿为主。")
    add_heading(doc, 2, "5.4 落地路径")
    add_para(doc, "第一阶段已完成六大模块 MVP 和云端部署；第二阶段优先围绕高校学生法律场景开展试用；第三阶段根据反馈完善校规边界、沟通模板和证据清单；第四阶段再向基层法律服务和机构合作扩展。")

    add_heading(doc, 1, "六、营销策略")
    add_heading(doc, 2, "6.1 市场定位")
    add_para(doc, "明鉴定位为“老百姓和学生身边的法律 AI助手”。品牌关键词是专业、简单、普惠和可执行。与传统数据库相比，它更强调结果可读和流程可操作；与通用 AI相比，它更强调法律场景结构化和本地依据约束。")
    add_heading(doc, 2, "6.2 线上推广策略")
    add_bullet(doc, "内容营销：围绕兼职押金不退、奖学金评定争议、宿舍矛盾、合同审查和起诉状生成制作短内容。")
    add_bullet(doc, "产品演示：通过 B站、抖音、公众号展示学生模块从表单填写到专题报告生成的过程。")
    add_bullet(doc, "搜索入口：围绕在线写起诉状、免费合同审查、学生兼职被骗怎么办等长尾关键词布局。")
    add_heading(doc, 2, "6.3 线下推广策略")
    add_para(doc, "以高校为重点试点场景，联合法学院、团委、学生工作部门和创新创业团队进行产品测试。学生模块适合用于普法活动、课程实践、比赛答辩和校园法律咨询分流。")
    add_heading(doc, 2, "6.4 合作伙伴策略")
    add_para(doc, "合作对象包括高校法学院、学生工作部门、中小型律所、基层法律服务机构和法律科技平台。早期合作目标不是快速商业化，而是获得真实反馈、优化模板和提高输出可靠性。")

    add_heading(doc, 1, "七、团队介绍")
    add_heading(doc, 2, "7.1 核心团队")
    add_para(doc, "团队信息可根据实际成员补充。建议按产品负责人、前端开发、后端开发、算法与知识库、法律内容审核、运营推广六类分工呈现。")
    add_heading(doc, 2, "7.2 团队优势")
    add_bullet(doc, "工程执行：已完成前端页面、后端服务、知识库、模型调用和部署。")
    add_bullet(doc, "场景聚焦：新增学生法律模块，贴合大学生创新大赛和高校真实需求。")
    add_bullet(doc, "迭代能力：项目已有可运行 MVP，可通过真实用户反馈持续优化。")

    add_heading(doc, 1, "八、财务分析")
    add_heading(doc, 2, "8.1 启动成本估算")
    add_table(doc, ["项目", "内容", "预估费用"], [
        ["云服务器", "阿里云 ECS 与带宽", "约 1,200 元/年"],
        ["模型调用", "国产模型 / OpenAI 兼容接口", "约 3,000-8,000 元/年"],
        ["域名与备案", "域名注册、备案与基础运维", "约 200 元/年"],
        ["知识库维护", "法条整理、风险模板和学生场景模板", "约 500-2,000 元/年"],
    ], [2100, 4700, 2200])
    add_heading(doc, 2, "8.2 年度运营预算")
    add_para(doc, "年度运营预算主要由服务器、模型调用、知识库维护、合规审核和小规模推广构成。学生模块上线后，预算重点应放在案例整理、模板优化和高校试点反馈，而不是大规模广告投放。")
    add_heading(doc, 2, "8.3 收入预测（3年）")
    add_table(doc, ["阶段", "第1年", "第2年", "第3年"], [
        ["核心目标", "验证留存", "验证付费", "验证规模化"],
        ["用户重点", "高校学生试用", "高校/基层协作", "机构化试点"],
        ["收入来源", "少量订阅和试点", "订阅+机构服务", "机构版+API"],
    ], [2100, 2300, 2300, 2300])
    add_heading(doc, 2, "8.4 盈亏平衡分析")
    add_para(doc, "第一年重点是验证产品价值和场景适配，形成稳定种子用户与试点案例。若学生模块能在高校场景形成可复制样板，后续专业版和机构版转化空间更清晰。")
    add_heading(doc, 2, "8.5 资金需求与用途")
    add_para(doc, "资金优先用于前后端功能完善、学生模块案例库建设、法律内容审核、数据安全和真实场景试点。")

    add_heading(doc, 1, "九、风险评估与对策")
    add_heading(doc, 2, "9.1 技术风险")
    add_para(doc, "AI 输出可能存在不准确或表达不严谨的问题。应对措施是使用本地知识库作为依据来源，输出免责声明，并保留人工审核和用户复核机制。")
    add_heading(doc, 2, "9.2 市场风险")
    add_para(doc, "C 端用户付费意愿有限，机构客户转化周期较长。项目应先用学生模块和高校试点验证真实需求，再逐步探索订阅和机构服务。")
    add_heading(doc, 2, "9.3 法律与合规风险")
    add_para(doc, "产品应明确定位为法律辅助工具，不替代律师服务、学校正式处理或司法机关判断。涉及重大权益的事项，应提示用户咨询专业律师或相关部门。")
    add_heading(doc, 2, "9.4 风险总结矩阵")
    add_table(doc, ["风险类别", "风险等级", "应对措施"], [
        ["AI 输出不准确", "高", "本地依据、免责声明、人工复核"],
        ["学生场景误用", "中高", "强调校内沟通、申诉流程和专业咨询"],
        ["数据隐私", "高", "最小化收集、加密传输、删除机制"],
        ["市场转化慢", "中", "先试点、再付费、再机构化"],
    ], [2300, 1800, 4900])

    add_heading(doc, 1, "十、发展规划")
    add_heading(doc, 2, "10.1 短期规划（3个月内）")
    add_numbered(doc, 1, "完善学生法律模块，补充校园管理、奖助学金、兼职被骗、实习劳动争议等模板。")
    add_numbered(doc, 2, "优化前端表单体验和移动端适配，提高学生用户填写效率。")
    add_numbered(doc, 3, "完善后端接口稳定性、历史记录和 PDF 文书导出能力。")
    add_heading(doc, 2, "10.2 中期规划（6-12个月）")
    add_numbered(doc, 1, "与高校法学院、团委或学生工作部门开展试点，收集真实反馈。")
    add_numbered(doc, 2, "建立学生法律案例库和校规边界模板库，提高输出可用性。")
    add_numbered(doc, 3, "形成高校普法活动和法学实训课程的标准演示方案。")
    add_heading(doc, 2, "10.3 长期愿景（1-3年）")
    add_numbered(doc, 1, "形成面向高校、基层法律服务和中小机构的 AI法律辅助平台。")
    add_numbered(doc, 2, "逐步开放 API 能力，与机构知识库、校内服务平台或法律科技平台对接。")
    add_heading(doc, 2, "10.4 里程碑时间线")
    add_table(doc, ["时间节点", "里程碑", "关键指标"], [
        ["2026年5月", "六大模块上线", "学生法律模块重点上线"],
        ["2026年8月", "学生模块完善", "完成种子用户反馈闭环"],
        ["2026年12月", "高校试点", "形成可演示样板场景"],
        ["2027年6月", "机构试点", "形成标准化交付方案"],
    ], [2300, 3300, 3400])

    add_heading(doc, 1, "十一、社会价值")
    add_heading(doc, 2, "11.1 降低法律服务门槛，推动法律普惠")
    add_para(doc, "明鉴通过低成本入口帮助用户完成第一轮事实整理、依据查询和材料生成，让基础法律服务更容易被普通人理解和使用。")
    add_heading(doc, 2, "11.2 促进法律知识普及，增强全民法治意识")
    add_para(doc, "系统通过自然语言检索、结构化表单和可读报告，把法律知识转化成用户可以理解和执行的步骤。")
    add_heading(doc, 2, "11.3 服务学生群体，补足校园法律教育入口")
    add_para(doc, "学生模块是项目最有辨识度的新增方向。它帮助学生在遇到校园或兼职问题时先明确事实、证据、校规边界和沟通路径，减少盲目对抗和无效投诉。")
    add_heading(doc, 2, "11.4 助力法治中国建设，缓解法律服务资源不均")
    add_para(doc, "AI 法律辅助工具不受物理网点限制，可以为基层法律服务、高校实践课程和普通公众提供标准化支持。")
    add_heading(doc, 2, "11.5 推动法律科技创新，促进产学研融合")
    add_para(doc, "项目把前端交互、后端服务、知识库、文书生成和学生场景结合起来，可作为高校法学教育、计算机课程和创新创业训练的交叉实践案例。")
    add_heading(doc, 2, "11.6 社会价值量化目标")
    add_table(doc, ["指标", "1年目标", "3年目标", "5年目标"], [
        ["累计服务人次", "3,000+", "100,000+", "500,000+"],
        ["学生专题分析", "500+", "20,000+", "100,000+"],
        ["免费文书生成", "300+", "10,000+", "50,000+"],
        ["试点场景", "5+ 个", "30+ 个", "100+ 个"],
        ["合作高校/机构", "3+", "20+", "80+"],
    ], [2300, 2200, 2200, 2200])
    add_heading(doc, 2, "11.7 数据与政策依据")
    add_bullet(doc, "司法部公开数据：全国法律服务机构、专业法律服务人员和公共法律服务实体平台规模持续扩大。")
    add_bullet(doc, "政策依据：《法治中国建设规划（2020-2025年）》提出推动现代科技与法治建设深度融合。")
    add_bullet(doc, "政策依据：公共法律服务体系建设相关政策提出覆盖城乡、便捷高效、均等普惠的现代公共法律服务体系目标。")
    add_bullet(doc, "政策依据：国务院《关于深入实施“人工智能+”行动的意见》提出推动人工智能与经济社会各领域深度融合。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
